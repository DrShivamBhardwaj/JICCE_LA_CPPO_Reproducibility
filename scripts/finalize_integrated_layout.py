from docx import Document
from docx.shared import Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
p=ROOT/'publication/updated_evidence/JICCE_LA_CPPO_Integrated_Review_Draft.docx'
d=Document(p)
for para in d.paragraphs:
 if para.text.startswith('[5] '):para.text='[5] T. Winter et al., “RPL: IPv6 routing protocol for low-power and lossy networks,” IETF, RFC 6550, 2012. DOI: 10.17487/RFC6550.'
 if para.text.startswith('['):
  import re
  para.text=re.sub(r'pp\. (\d{3,7})(?=[,.])',r'Art. no. \1',para.text)
  from docx.shared import Pt
  para.paragraph_format.space_after=Pt(0)
  for run in para.runs:run.font.name='Times New Roman';run.font.size=Pt(9.5)
for table in d.tables:
 table.autofit=False
 grid=table._tbl.tblGrid
 widths=[col.w for col in grid.gridCol_lst]
 total=sum(widths)
 scaled=[int(Cm(16)*w/total) for w in widths]
 for col,width in zip(grid.gridCol_lst,scaled):col.w=width
 for row in table.rows:
  for cell,width in zip(row.cells,scaled):
   cell.width=width
   pr=cell._tc.get_or_add_tcPr()
   for shading in list(pr.findall(qn('w:shd'))):pr.remove(shading)
   for p in cell.paragraphs:
    for r in p.runs:
     from docx.shared import RGBColor
     r.font.color.rgb=RGBColor(0,0,0)
 pr=table._tbl.tblPr
 width=pr.find(qn('w:tblW'))
 if width is not None:width.set(qn('w:type'),'dxa');width.set(qn('w:w'),'9071')
 indent=pr.find(qn('w:tblInd'))
 if indent is not None:indent.set(qn('w:w'),'0')
from PIL import Image
from io import BytesIO
for shape in d.inline_shapes:
 rid=shape._inline.graphic.graphicData.pic.blipFill.blip.embed
 part=d.part.related_parts[rid]
 if str(part.partname).endswith('image7.png'):part._blob=(ROOT/'publication/updated_evidence/figures/fig7.png').read_bytes()
 image=Image.open(BytesIO(part.blob))
 width=min(Cm(16),Cm(image.width/900*2.54))
 shape.width=int(width);shape.height=int(width*image.height/image.width)
for para in d.paragraphs:
 if para.text.startswith('Metaheuristic clustering and learned route selection') and 'Table I summarizes' not in para.text:para.add_run(' Table I summarizes the methodological positioning.')
 if para.text.startswith('Traffic is divided into three classes') and 'Tables II and III' not in para.text:para.add_run(' Tables II and III summarize the configuration and threat-model visibility.')
 if para.text.startswith('Figures 3 and 4 show seedwise') and 'Table IV reports' not in para.text:para.add_run(' Table IV reports the paired inferential results.')
 if para.text.startswith('Attention passed both non-inferiority checks') and 'Table V reports' not in para.text:para.add_run(' Table V reports the estimates and frozen margins.')
 if para.text.startswith('The nominal comparison family consists') and 'Table IX summarizes' not in para.text:para.add_run(' Table IX summarizes the nominal and transfer outcomes.')
# Preserve native notation; no math paragraphs are rewritten.
assert len(d.element.xpath('.//m:oMath'))==65
assert len(d.inline_shapes)==7
d.core_properties.title=d.paragraphs[0].text
d.core_properties.subject='JICCE redesigned manuscript review draft'
d.core_properties.comments='Includes exploratory localhost UDP diagnostics with explicit deployment limitations.'
d.save(ROOT/'publication/updated_evidence/JICCE_LA_CPPO_Integrated_Review_Draft.docx')
