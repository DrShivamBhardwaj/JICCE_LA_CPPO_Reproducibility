"""Place all exhibits in the body while preserving native equations/tables."""
import json,re
from copy import deepcopy
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt,Cm

def text_of(el):return ''.join(n.text or '' for n in el.iter(qn('w:t')))
def apply_revision(doc, root):
    data=json.loads((root/'publication/array/prose_revision.json').read_text());count=0
    for entry in data['paragraphs']:
        matches=[p for p in doc.paragraphs if p.text==entry['match']]
        if len(matches)!=1:raise ValueError('Prose revision match missing/duplicated: '+entry['match'][:70])
        p=matches[0];parts=entry['replacement'];p.text=parts[-1]
        for part in parts[:-1]:p.insert_paragraph_before(part,style=p.style)
        count+=1
    return count

def place_exhibits(doc):
    body=doc._element.body
    def find(prefix):
        found=[p._p for p in doc.paragraphs if p.text.startswith(prefix)]
        if len(found)!=1:raise ValueError('Placement anchor not unique: '+prefix)
        return found[0]
    groups={}
    for n in range(1,7):
        cap=find(f'Table {n}. ');node=cap.getnext()
        while node is not None and node.tag!=qn('w:tbl'):node=node.getnext()
        if node is None:raise ValueError('Missing table '+str(n))
        group=[cap,node]
        note=node.getnext()
        if note is not None and text_of(note).startswith(('Values are censored','Means over 30 fixtures')):group.append(note)
        groups['table'+str(n)]=group
    for n in range(1,9):
        cap=find(f'Fig. {n}. ');img=cap.getprevious()
        while img is not None and img.find('.//'+qn('w:drawing')) is None:img=img.getprevious()
        if img is None:raise ValueError('Missing image '+str(n))
        groups['fig'+str(n)]=[img,cap]
    start=find('Algorithm 1 Observer');end=find('Table 1. ');nodes=[];el=start
    while el is not None and el is not end:
        if text_of(el).strip():nodes.append(el)
        el=el.getnext()
    groups['algorithm1']=nodes
    anchors={
        'table1':'Recent studies combine trust',
        'table2':'Runs last 4 s.',
        'table3':'Table 3 sets out the feedback controls.',
        'fig1':'In this benchmark, only the source updates service debts',
        'algorithm1':'Algorithm 1 gives the observer procedure',
        'fig2':'Algorithm 1 gives the observer procedure',
        'fig3':'Thirty fixtures crossed with five settings',
        'fig4':'Under blackholes, censored passive attention delivered',
        'table4':'Table 4 reports the remaining blackhole contrasts.',
        'fig5':'Table 4 reports the remaining blackhole contrasts.',
        'fig6':'Figure 6 separates total energy',
        'table5':'All three blackhole margins against frozen prior passed.',
        'fig7':'The authenticated branch uses the message path',
        'table6':'Table 6 and Fig. 8 summarize all 180 runs.',
        'fig8':'Table 6 and Fig. 8 summarize all 180 runs.'}
    after={}
    for key,prefix in anchors.items():
        anchor=after.get(prefix)
        if anchor is None:anchor=find(prefix)
        for node in groups[key]:anchor.addnext(node);anchor=node
        after[prefix]=anchor
    # Delete old exhibit-only page breaks and empty spacer paragraphs.
    for p in list(doc.paragraphs):
        if not p.text.strip() and p._p.find('.//'+qn('w:drawing')) is None and p._p.find('.//'+qn('m:oMath')) is None:
            if p._p.getparent() is not None:p._p.getparent().remove(p._p)
    for p in doc.paragraphs:
        p.paragraph_format.page_break_before=False
        p.paragraph_format.widow_control=True
        p.paragraph_format.space_after=Pt(6)
        p.paragraph_format.line_spacing=1.15
        if p.style.name.startswith('Heading'):p.paragraph_format.keep_with_next=True
        elif p.text.startswith('Fig. '):
            p.paragraph_format.keep_with_next=False;p.paragraph_format.keep_together=True
            for run in p.runs:run.font.size=Pt(10)
        elif p.text.startswith('Table '):p.paragraph_format.keep_with_next=True
        elif p._p.find('.//'+qn('w:drawing')) is not None:
            p.paragraph_format.keep_with_next=True;p.paragraph_format.keep_together=True
        else:p.paragraph_format.keep_with_next=False
    for shape in doc.inline_shapes:
        if shape.height>Cm(17.5):
            ratio=Cm(17.5)/shape.height;shape.width=int(shape.width*ratio);shape.height=Cm(17.5)
    for table_index,table in enumerate(doc.tables):
        table.autofit=False
        for i,row in enumerate(table.rows):
            pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
            if i==0:pr.append(OxmlElement('w:tblHeader'))
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next=i<len(table.rows)-1;p.paragraph_format.space_after=Pt(1 if table_index==5 else 4);p.paragraph_format.space_before=Pt(1 if table_index==5 else 4);p.paragraph_format.line_spacing=1.05
                    for r in p.runs:r.font.size=Pt(9)
    # Keep equation introductions, displays and variable definitions together.
    ps=doc.paragraphs
    for i,p in enumerate(ps):
        if i+1<len(ps) and ps[i+1]._p.find('.//'+qn('m:oMath')) is not None:p.paragraph_format.keep_with_next=True
        if p._p.find('.//'+qn('m:oMath')) is not None:p.paragraph_format.keep_with_next=True
        if p.text.startswith('•') and i+1<len(ps) and ps[i+1].text.startswith('•'):p.paragraph_format.keep_with_next=True
    # Keep the short algorithm together without retaining all former appendices.
    for p in doc.paragraphs:
        if p._p in groups['algorithm1']:
            p.paragraph_format.keep_with_next=p._p is not groups['algorithm1'][-1]
            p.paragraph_format.space_after=Pt(3)
            p.paragraph_format.line_spacing=1.05
            for r in p.runs:r.font.size=Pt(10)
    return {'figures_in_body':8,'tables_in_body':6,'algorithm_in_body':1}
