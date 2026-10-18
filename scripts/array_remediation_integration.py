"""Integrate only completed, independently rebound remediation evidence."""
from pathlib import Path
from copy import deepcopy
import json
from docx.shared import Cm,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from verify_remediation_archive import verify_dataset
TITLE='Controlled Feedback Comparisons and Adaptation Limits in Authenticated Wireless Sensor Routing'
LABELS={'negative_main':'Negative update','proof_main':'Proof generation','interaction':'Interaction'}

def primary_rows(a):
    rows=[]
    for timing in [0,1]:
        for key in ['negative_main','proof_main','interaction']:
            x=next(x for x in a['contrasts'] if (x['endpoint'],x['timing'],x['contrast'])==('post_delivery',timing,key))
            rows.append(['Native' if timing==0 else 'Scheduled',LABELS[key],f"{100*x['mean_difference']:+.2f}",f"[{100*x['ci95'][0]:+.2f}, {100*x['ci95'][1]:+.2f}]",f"{x['p_holm']:.4g}"])
    return rows

def effect(x):
    return f"{100*x['mean_difference']:+.2f} percentage points (95% pointwise interval {100*x['ci95'][0]:+.2f} to {100*x['ci95'][1]:+.2f}; Holm-adjusted p={x['p_holm']:.4g})"

def apply(doc,root):
    root=Path(root);out=root/'publication/array';dataset=root/'outputs/remediation_v1'
    if any(p.text=='5.10 Archive diagnostics' for p in doc.paragraphs):raise ValueError('Remediation already integrated')
    integrity=verify_dataset(dataset,root/'experiments/remediation_v1/protocol.json')
    a=json.loads((dataset/'analysis.json').read_text());d=json.loads((dataset/'archive_diagnostics/summary.json').read_text());mcu=json.loads((root/'outputs/mcu_deadline_v1/analysis.json').read_text())
    if a['runs']!=1440 or d['archived_runs']!=600 or d['matched_decisions']!=11927 or d['risk_different_decisions']!=7023 or d['action_disagreements']!=0 or d['offline_interventions']!=4800 or d['offline_action_changes']!=1010:raise ValueError('Unreviewed headline inputs')
    paths=json.loads((dataset/'archive_diagnostics/path_summary.json').read_text())
    if paths['matched_decisions']!=11927 or paths['logits_differ']!=7023 or paths['feedback_rows']!=143928:raise ValueError('Unreviewed recorded-logit paths')
    misses=sum(x['signed_deadline_misses'] for x in mcu['rows'])
    if misses!=68 or sum(x['signed_receipts'] for x in mcu['rows'])!=200:raise ValueError('Unreviewed MCU outcomes')
    neg=next(x for x in a['contrasts'] if (x['endpoint'],x['timing'],x['contrast'])==('post_delivery',1,'negative_main'))
    proof=next(x for x in a['contrasts'] if (x['endpoint'],x['timing'],x['contrast'])==('post_delivery',1,'proof_main'))
    abstract=('Forwarding silence can reflect loss rather than misconduct, but censoring it also prevents missing observations from reversing earlier positive evidence. We examine this trade-off in a separately frozen 1,440-run Cooja factorial study and retain the earlier 750-, 180-, 240- and 600-run experiments. Twelve perturbed geometry blocks combine three retained checkpoints, two radio seeds per block, relay-identity permutations, both attacker positions, benign controls and two timing settings. Four matched arms vary missing-evidence penalties and proof generation while sharing ACK eligibility, expiry and positive-credit handling; an active-proof reference retains its legacy evidence semantics. With scheduled processing, enabling negative updates changes post-onset delivery by '+effect(neg)+'. Enabling proof generation changes it by '+effect(proof)+'. Archive diagnostics find 7,023 risk-coordinate differences across 11,927 matched censored–frozen decisions without a relay-choice change; 1,010 of 4,800 offline risk interventions change actions. Compiler-only O3 MCU execution still misses 68 of 200 original signed-receipt deadlines. The contribution is auditable attribution of feedback and traffic interventions within a declared fixture family, rather than a new learning algorithm or a universal switching threshold. Public fixture keys, narrow geometry coverage and illustrative scheduled delays restrict security and deployment inference. Physical RF, measured energy and production device protections remain unvalidated.')
    if len(abstract.split())>250:raise ValueError('Abstract exceeds 250 words')
    changes=[]
    def replace(p,new):
        if p.text!=new:changes.append({'before':p.text,'after':new});p.text=new
    for p in doc.paragraphs:
        if p.style.name=='Title':replace(p,TITLE)
    abstract_p=doc.paragraphs[next(i for i,p in enumerate(doc.paragraphs) if p.text=='Abstract')+1];replace(abstract_p,abstract)
    p=next(p for p in doc.paragraphs if p.text.startswith('The original prespecified decision combines'))
    replace(p,'The original prespecified decision combines a blackhole delivery comparison against active proofs with delivery, critical-service and energy-efficiency margins against frozen-prior attention. Sections 5.10–5.11 add exploratory diagnostics and a separately frozen matched-factorial experiment. The latter isolates missing-evidence penalties and proof generation under common ACK, expiry and late-positive rules, while counterbalancing relay identity and attacker position. These are controlled comparisons and implementation checks; they do not establish a new learning algorithm, a universal operating boundary or production security.')
    p=next(p for p in doc.paragraphs if p.text.startswith('The benchmark exposes a failure condition rather than estimating its onset.'))
    replace(p,'The historical Cooja study exposes a failure condition, while the remediation study varies onset, channel condition, geometry and retained checkpoint within a declared small fixture family. Two onset settings and three checkpoints do not identify a universal change-point or a safe passive-feedback operating region. The balanced block design also does not fully cross every geometry with every checkpoint/onset/channel setting. An optimal probing policy remains a prospective mechanism requiring its own frozen evaluation.')
    p=next(p for p in doc.paragraphs if p.text.startswith('The evidence supports a conditional feedback comparison.'))
    replace(p,p.text+' Section 5.11 removes the stated ACK/expiry/late-positive differences within four new matched arms. Its factorial effects apply to that intervention and fixture family; they do not retroactively identify the complete cause of the original 750-run gain.')
    p=next(p for p in doc.paragraphs if p.text.startswith('The benchmark identifies a conditional benefit and a consequential transfer failure.'))
    replace(p,'Matched feedback controls now separate the missing-evidence update rule from proof generation in 1,440 Cooja runs over twelve perturbed-geometry blocks. In the scheduled setting, their post-onset delivery effects are '+effect(neg)+' for negative updates and '+effect(proof)+' for proof generation. These estimates complement the retained conditional 4.67-point gain and later transfer failure; they do not establish a universally preferable mechanism. Historical policy comparisons supply no established attention advantage, and the remediation experiment does not retest architectural superiority.')
    p=next(p for p in doc.paragraphs if p.text.startswith('Censored feedback with proof generation enabled delivers 32.83%'))
    replace(p,'Exploratory reconstruction explains the censored–frozen equality in the archived setting: recorded risk changes did not alter selected relays, although larger offline risk interventions could. Failure under the tested conditions therefore cannot be equated with complete actor insensitivity. A compiler-only O3 pilot retains 68/200 signed-receipt deadline misses. The study supplies neither a universal switching threshold nor a production hardware validation. Physical RF, measured energy, protected device keys and reboot-persistent MCU state remain unvalidated.')
    # Add O3 execution outcome without modifying the retained six-row table.
    p=next(p for p in doc.paragraphs if p.text.startswith('Fresh pairwise keys and a signing seed are provisioned once'))
    text='A separately frozen compiler-only O3 pilot repeats one hundred packets in each topology with the same Q22 actor, signature-verification order and original deadlines. Both runs deliver all offered packets and source-verified Ed25519 receipts, and all five hundred decisions match the frozen reference. Signed delays are 331–344 ms in the chain and 239–247 ms in the choice geometry; each run still misses thirty-four deadlines. Seven authenticated negative controls pass per run. The Linux Renode run and historical Mac O2 run are descriptive diagnostics, not paired statistical replicates. All sixty-eight misses are in the 90 ms critical class; neither of the other classes misses its deadline. UART stage records show median sink signing of 30 ms and median signature verification of 82 ms at the source and forwarding relays; per-call HMAC medians are 1 ms. These cryptographic stages dominate the emulated receipt path. The stage records include control operations and are descriptive emulator timings, not independent packet replicates or physical measurements. Compiler optimization therefore does not resolve the original deadline requirement. Public rechecking verifies all two hundred sink-generated signatures; private hop-tag rechecking requires the intentionally unexported ephemeral fixture keys.'
    first,second=text.split(' All sixty-eight',1)
    added=[first,'All sixty-eight'+second]
    for part in added:
        inserted=p.insert_paragraph_before(part)
        if p._p.pPr is not None:inserted._p.insert(0,deepcopy(p._p.pPr))
    anchor=next(p for p in doc.paragraphs if p.text=='6. Discussion');template=next(p for p in doc.paragraphs if p.text.startswith('The simulated network has'))._p.pPr
    def para(text,style=None):
        p=doc.add_paragraph(text,style=style)
        if style is None and template is not None:
            if p._p.pPr is not None:p._p.remove(p._p.pPr)
            p._p.insert(0,deepcopy(template))
        anchor._p.addprevious(p._p);added.append(text)
        if style or text.startswith('Table '):p.paragraph_format.keep_with_next=True
        return p
    para('5.10 Archive diagnostics','Heading 2')
    para('The archived censored and frozen-prior captures are compared post hoc by simulation seed, attack/timing stratum, observer and packet identity. All 11,927 decisions across 120 paired runs match; no decision is unmatched. Risk coordinates and recorded valid logits differ in 7,023 decisions, while non-risk coordinates, candidate ordering and selected relays remain identical. The maximum paired valid-logit difference is 0.15597. Effective-count trajectories preserve all 143,928 archived feedback events. Figure 13 separates these descriptive counts from deadline outcomes. Equality of delivery is therefore accompanied by equality of recorded actions in this archive, rather than being used to infer it.')
    para('A deterministic grid intervenes on one candidate’s two risk coordinates at a time in 120 evenly spaced multi-choice source observations per attack/timing stratum. The five effective-count pairs are (1,1), (5,1), (20,1), (1,5) and (1,20); the retained candidate set and other coordinates are held fixed. The compiled Q22 actor changes its action in 1,010 of 4,800 interventions, with maximum logit change 0.18371. These are offline actor-sensitivity checks, not closed-loop delivery results, independent replicates or a test of candidate preselection. They show that the observed positive-count changes were insufficient to change choices, not that every risk input is ignored.')
    para('For benign censored runs with scheduled delays, 600 critical packets include thirty timely deliveries, 563 delivered-but-late packets and seven undelivered packets. The median sink application delay is 131 ms and source signed-receipt delay 382 ms over their respective received packets across all classes. Thus, most critical failures are lateness rather than loss. Sink arrival, application processing and reverse receipt timing remain distinct endpoints; high delivery does not establish the original 90 ms service target. Antecedent-event checks pass on all 600 archived runs, binding emitted feedback to prior received frames, processing readiness, own MAC outcomes and packet identities. They complement wire/ledger/policy audits rather than proving an independent RF model.')
    fig=para('');fig.paragraph_format.keep_with_next=True;fig.add_run().add_picture(str(out/'remediation_figures/fig13_feedback_diagnostics.png'),width=Cm(16));para('Fig. 13. Feedback and deadline diagnostics.')
    para('5.11 Matched feedback and counterbalanced transfer','Heading 2')
    para('The new protocol was committed before new radio outcomes as b24418008b318aa4a082e0a666c4a2196b5c994c; pre-primary-outcome amendment 40a933ee06b0dd1521f7de36d4b215f9f8d8b9d4 records GCC portability corrections after an excluded compile-only preflight failure. Twelve independently perturbed four-node geometries use three retained checkpoints. Each checkpoint covers onset packet 10 or 30 crossed with UDGM RX probability 0.9 or 1. Two radio seeds per block cross both relay-ID permutations, benign behavior or either physical attacker position, both timing settings and five arms: 1,440 runs. Perturbations are uniform within ±1 m around each original relay position. Source and sink remain fixed; the geometry family is narrow and the block factors are not fully crossed within each geometry.')
    para('Four arms cross a missing-evidence failure increment with proof generation. Algorithm 1 and Fig. 2 retain the censored lifecycle; the negative-update arms add the declared failure increment at expiry. Figures 1 and 7 continue to describe their original simulation and host-transport stages, not a production deployment. All four require the observer’s own ACK, expire 150 ms after it, apply the same incremental clipped positive weight and admit late positive evidence until slot reuse. Missing-evidence penalties use the cached expected-link weight once at expiry. A late positive observation adds positive evidence without retracting that earlier penalty. Proof frames consume airtime but do not enter these four ledgers; route-dependent realized proof counts can differ. The fifth arm retains legacy active-proof eligibility and expiry as a separate reference, with the shared stricter proof source/session/pending-route check. Public fixture keys exercise serialization and validation, not deployment secrecy. Native execution and overlapping per-packet scheduled delays retain the scope of Section 5.9.')
    para('The remediation questions were motivated by the archived failure; they are exploratory relative to the original study, with their own analysis prespecified before the new runs. Primary inference within this new family uses post-onset unique sink delivery in attacked runs. Each geometry block averages the two radio seeds, two identity permutations and two physical attacker positions within timing and arm. Negative-update and proof-generation main effects average their two corresponding simple effects; the interaction is the difference in differences. Pointwise 95% intervals use 20,000 paired geometry-block bootstrap resamples. Two-sided exhaustive sign-flip tests enumerate 4,096 sign assignments over twelve block differences, conditional on symmetric block effects; this is not a claim of randomized treatment allocation. Holm adjustment covers six primary contrasts, three per timing setting. Whole-run critical failures and total transmit-state seconds have separate six-contrast secondary families. No equivalence or noninferiority claim is made. Table 12 and Fig. 14 report the primary effects.')
    para('Table 12. Matched delivery effects.')
    table=doc.add_table(rows=1,cols=5);table.autofit=False
    for i,row in enumerate([['Timing','Contrast','Effect (pp)','95% interval (pp)','Holm p']]+primary_rows(a)):
        cells=table.rows[0].cells if not i else table.add_row().cells
        for cell,text,width in zip(cells,row,[2.4,3.5,2.3,4.2,2.0]):
            cell.width=Cm(width);cell.text=text
            for p in cell.paragraphs:
                p.paragraph_format.space_before=p.paragraph_format.space_after=Pt(3)
                for run in p.runs:run.font.name='Times New Roman';run.font.size=Pt(9);run.bold=i==0
        table.rows[i]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    borders=OxmlElement('w:tblBorders')
    for side in ['top','bottom','insideH','left','right','insideV']:
        x=OxmlElement('w:'+side);x.set(qn('w:val'),'single' if side in ['top','bottom','insideH'] else 'nil');x.set(qn('w:sz'),'4');borders.append(x)
    table._tbl.tblPr.append(borders);anchor._p.addprevious(table._tbl)
    para('Effects use percentage points and equal geometry-block weight; packets are not independent replicates. Native means no added compute delay, not calibrated zero MCU processing cost. The intervention estimates include downstream routing and traffic consequences. The active reference is not part of the matched factorial family, and its additional comparisons remain exploratory. All source captures, compressed native images, case settings and checkpoint identities are rebound by hash before this manuscript integration; historical statistical families remain unchanged.')
    fig=para('');fig.paragraph_format.keep_with_next=True;fig.add_run().add_picture(str(out/'remediation_figures/fig14_matched_feedback_effects.png'),width=Cm(16));para('Fig. 14. Matched feedback effects.')
    secondary=json.loads((dataset/'secondary_summary.json').read_text())
    delayed=[g for g in secondary['groups'] if g['timing']==1]
    labels={'censored':'censored','censored_proof':'censored with proofs','negative':'negative update','negative_proof':'negative update with proofs','active_reference':'active reference'}
    recovery='; '.join(f"{labels[g['arm']]} {g['recovered_within_20']}/{g['attacked_runs']}" for g in delayed)
    censoring='; '.join(f"{labels[g['arm']]} {g['right_censored']}/{g['attacked_runs']}" for g in delayed)
    para('Under scheduled processing, the numbers of attacked run settings meeting the fixed twenty-opportunity recovery criterion are '+recovery+'. Right-censored settings over the full available window are '+censoring+'. These are descriptive tallies over correlated seed, permutation and attacker settings, not additional independent experimental units. The retained secondary summary reports conditional confirmation latencies alongside censoring, cumulative post-onset losses, relay switches and proof frames/bytes. Feedback trajectories preserve effective counts and their heuristic risk coordinates; no observed recovery fraction is used to tune the policy or select an onset threshold.')
    para('Adaptation is assessed as the first three consecutive source choices avoiding the declared attacker, with right censoring when no such sequence occurs. Recovery within twenty packet opportunities uses the same fixed observation window for both onset settings; a single safe choice is not sustained recovery. This is a routing-choice diagnostic rather than proof of subsequent delivery. The budget fixes twelve geometry blocks; it is not a power justification for declaring small effects absent. The small balanced fixture family, shared four-node graph, fixed source/sink, three checkpoints, public keys and nominal energy/queue telemetry limit transfer. Hop-authenticated relay claims also do not establish truthful forwarding by a malicious key holder, and the runtime router cannot access the independent auditor’s global trace.')
    (out/'remediation_edits.json').write_text(json.dumps({'classification':'Integration of audited prospective intervention and explicitly exploratory archive diagnostics','replacements':changes,'added_paragraphs':added,'new_table_rows':primary_rows(a),'integrity':integrity},indent=2,ensure_ascii=False)+'\n')
    return {'remediation_runs':1440,'figures':14,'tables':12,'references':25,'abstract_words':len(abstract.split()),'original_signed_deadline_complete':False}
