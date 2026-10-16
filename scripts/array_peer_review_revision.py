"""Review-driven clarifications without changing frozen evidence."""
from docx.shared import Pt
TITLE='A Reproducible Benchmark of Receipt Assisted Passive Feedback and Authenticated Routing in Wireless Sensor Networks'
def transform(text):
 if text.startswith('Receipt Assisted Feedback and Authenticated Policy Transfer'):return TITLE
 if text.startswith('Routing feedback uses airtime,'):
  text=text.replace('Compiled radio-interface firmware then runs on four emulated Cortex-M3 boards, with all three host-policy controls delivering ten of ten packets.','Compiled UART-to-radio bridge firmware runs on four emulated Cortex-M3 boards; neural inference and cryptography execute on the host PC. Each host-policy control delivers ten of ten packets.')
 if text.startswith('Figure 3 traces the experimental stages.'):
  text=text.replace('Its PRISMA-inspired layout accounts for runs and design decisions; it is not a systematic-review screening diagram.','The experimental protocol pipeline records source freezing, the technical pilot and all retained runs.')
 if text.startswith('The evaluation addresses three questions.'):
  return 'The benchmark tests blackhole delivery against active proofs under the prespecified multiple-testing rule. Delivery, critical-service and energy-efficiency margins against frozen-prior attention assess whether the feedback mechanism can be replaced. Naive-passive and capacity-matched controls test attribution to censoring or policy architecture. These comparisons distinguish a useful replacement from a unique algorithmic advantage.'
 if text.startswith('Routing feedback uses airtime,'):
  text=text.replace('The evidence supports the tested feedback replacement', 'Post-review 60–300 s stress runs expose early routing-energy exhaustion and increasing service debt. The evidence supports the tested feedback replacement')
 if text.startswith('Receipt-assisted passive feedback passed'):
  text=text.replace('These results support', 'The 60–300 s extension exposes routing-energy exhaustion and increasing source debt rather than sustained operation. These results support')
 return text.replace('1.1 Contribution and research questions','1.1 Benchmark contribution').replace('6.2 High novelty opportunities','6.2 Future experiments').replace('5.2 Frozen policy transfer protocol','5.2 Frozen host policy transport protocol').replace('5.4 Virtual serial devices and firmware boundary','5.4 Host controlled radio bridge firmware')
ADDITIONS=[
('Shared encoders map context','The coefficients in Eq. (1) are fixed hand-set preselection constants, rather than learned parameters. The retained source records their values but supplies no systematic tuning record. They are treated as benchmark settings. Equation (2) divides the query–key product by the square root of 96 in both the native equation and the PyTorch implementation; multiplication by that factor is not implemented.'),
('The policy retains its existing evidence-risk','Equation (4) applies clipping only when the uncredited evidence increment is positive. Its static floor is a heuristic and can alter the total assigned credit when one observation is split into successive increments. It is not a Bayesian count update. In Eq. (5), 1.645 is likewise a fixed score multiplier, not a calibrated 95% coverage guarantee or a Beta quantile. Replacing either rule changes the feedback mechanism and requires a separate experiment.'),
('Thirty fixtures crossed','The four-second horizon isolates a burst and its feedback. Positions are available at startup and telemetry begins immediately; the model does not simulate RPL or CTP parent discovery. This design therefore evaluates a transient workload rather than routing convergence or steady operation.'),
('Figure 6 separates total energy','Duty cycling changes the cost of passive observation. As an accounting example, an extra 150 ms listening window instead of radio-off time costs 3 × (0.0188 − 0.000020) × 0.15 = 0.008451 J under the assumed currents. Overlapping windows and existing awake periods reduce the incremental cost. This calculation is not a duty-cycled MAC experiment. Such a comparison must charge channel checks, rendezvous, observation windows and proof retransmissions to all arms; the always-listening results do not predict which mechanism would use less energy.'),
('The transfer runs show','The −1.78-point attention–greedy difference under selective dropping has an interval that crosses zero, so it does not establish greedy superiority. LQI replacing the training RSSI proxy, raw queue counts and receipt-only forwarding evidence are plausible sources of observation shift; their causal contributions were not isolated. Greedy is the economical starting point for deployment assessment because it avoids neural weights and tensor inference. It still requires neighbor state, authentication and packet processing.'),
("Renode's deterministic range",'The full attention actor–critic stores 63,368 FP32 parameters, or 253,472 bytes (247.53 KiB). That count includes critic parameters that an actor-only implementation could omit. For comparison, the CC2538SF53 variant provides 512 KB flash and 32 KB RAM [22]. Flash-resident weights need not all occupy SRAM, but flash capacity alone says nothing about activation buffers, stack and protocol memory. An embedded build, linker map and peak-memory measurement are needed before claiming feasibility. The present MCU firmware handles UART parsing, frame bounds, radio FIFO access and MAC acknowledgments; it executes neither neural inference nor HMAC/Ed25519.'),
('The holdout fixes passive strength','The 150 ms evidence timeout includes contention, retries and forwarding. Interference can therefore delay a valid receipt past expiry. Censoring avoids assigning failure solely from silence and allows a later receipt to upgrade the record; it does not establish an optimal timeout. The simulator also uses a shared time base. The corrected host timestamp sampling error is different from device clock offset or drift, which remain untested.')]
def apply(doc):
 for p in doc.paragraphs:
  s=transform(p.text)
  if s!=p.text:p.text=s
  if p.text.startswith('Algorithm 1 Observer'):
   p.text=p.text.replace('Algorithm 1 Observer','Algorithm 1: Observer');p.alignment=0
 for anchor,text in ADDITIONS:
  p=next(p for p in doc.paragraphs if p.text.startswith(anchor));p=p.insert_paragraph_before(text)
  p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(6);p.paragraph_format.widow_control=True
 from docx.oxml import OxmlElement
 from docx.text.paragraph import Paragraph
 ref=next(p for p in doc.paragraphs if p.text.startswith('[21]'));el=OxmlElement('w:p');ref._p.addnext(el);p=Paragraph(el,ref._parent)
 p.text='[22] Texas Instruments, “CC2538 Powerful Wireless Microcontroller System-On-Chip,” datasheet SWRS096D, rev. D. Available: https://www.ti.com/lit/ds/symlink/cc2538.pdf. Accessed October 16, 2026.'
 p.paragraph_format.line_spacing=1.15
