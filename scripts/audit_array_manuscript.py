"""Check evidence, equations, tables, authors and inline exhibit placement."""
from pathlib import Path
import json,re,hashlib,zipfile
from lxml import etree
from docx import Document
from docx.oxml.ns import qn
from array_peer_review_revision import transform
from array_failure_boundary_revision import transform_text
R=Path(__file__).resolve().parents[1];P=R/'publication/array';d=Document(P/'Array_Manuscript.docx');base=Document(R/'publication/redesigned/JICCE_Redesigned_Manuscript.docx')
old=base._element.xpath('//m:oMath');new=d._element.xpath('//m:oMath')
assert len(old)==len(new)==10
assert [etree.tostring(e) for e in old]==[etree.tostring(e) for e in new],'Native equations changed'
remediated=(P/'remediation_edits.json').exists()
expected_figures=14 if remediated else 12;expected_tables=12 if remediated else 11
assert len(d.inline_shapes)==expected_figures and len(d.tables)==expected_tables
for a,b in zip(base.tables,d.tables[:5]):
 assert [[c.text for c in row.cells] for row in a.rows]==[[c.text for c in row.cells] for row in b.rows],'Frozen table values changed'
anal=json.loads((R/'outputs/auth_policy/analysis.json').read_text());actual=d.tables[5]
for i,(setting,arm) in enumerate((s,a) for s in ['benign','selective_drop'] for a in ['greedy','mlp','attention']):
 z=next(r for r in anal['arm_means'] if r['setting']==setting and r['arm']==arm)
 expected=[('Benign' if setting=='benign' else 'Dropping')+' / '+arm,f"{100*z['delivery']:.2f}",f"{100*z['verified_receipts']:.2f}",f"{100*z['critical_violation']:.2f}",f"{z['total_radio_j']:.4f}"]
 assert [c.text for c in actual.rows[i+1].cells]==expected
extension=json.loads((R/'outputs/review_extensions/analysis.json').read_text())
for i,z in enumerate(extension['long_groups']):
 label={'active_proof':'Active proofs','censored_passive':'Censored passive','frozen_prior':'Frozen prior'}[z['arm']]
 expected=[f"{z['duration_s']} s / {label}",str(z['packets_per_run']),f"{z['delivery_mean_percent']:.2f}",f"{z['critical_failure_mean_percent']:.2f}",f"{z['source_debt_max']:.2f}"]
 assert [c.text for c in d.tables[6].rows[i+1].cells]==expected,'Stress table mismatch'
factorial=json.loads((R/'outputs/validation_20261018/factorial_analysis.json').read_text())
for i,z in enumerate(factorial['arm_means']):
 expected=[f"{100*z['duty_fraction']:.0f}%",'Censored' if z['arm'].startswith('censored') else 'Naive','On' if z['arm'].endswith('with_proof') else 'Off',f"{z['delivery_percent']:.2f}",f"{z['critical_violation_percent']:.2f}",f"{z['radio_energy_j']:.6f}"]
 assert [c.text for c in d.tables[7].rows[i+1].cells]==expected
assert [c.text for c in d.tables[8].rows[4].cells]==['Text and read-only data','341,600 bytes']
assert [c.text for c in d.tables[8].rows[5].cells]==['Observed stack watermark','3,080 bytes']
text='\n'.join(p.text for p in d.paragraphs);refs=next(i for i,p in enumerate(d.paragraphs) if p.text=='References')
assert len([p for p in d.paragraphs[refs+1:] if re.match(r'^\[\d+\]',p.text)])==25
figs=[(i,int(m.group(1))) for i,p in enumerate(d.paragraphs) if (m:=re.match(r'^Fig\. (\d+)\. ',p.text))]
tables=[(i,int(m.group(1))) for i,p in enumerate(d.paragraphs) if (m:=re.match(r'^Table (\d+)\. ',p.text))]
assert [n for _,n in figs]==list(range(1,expected_figures+1)) and [n for _,n in tables]==list(range(1,expected_tables+1))
assert all(i<refs for i,_ in figs+tables),'An exhibit is still appended after references'
for i,n in figs:
 prior=' '.join(p.text for p in d.paragraphs[:i]);assert re.search(r'(?:Fig\.|Figure)\s+'+str(n)+r'\b',prior),f'Figure {n} lacks a prior discussion reference'
for i,n in tables:
 prior=' '.join(p.text for p in d.paragraphs[:i]);assert re.search(r'Table\s+'+str(n)+r'\b',prior),f'Table {n} lacks a prior discussion reference'
revision=json.loads((P/'prose_revision.json').read_text())
editorial={x['before']:x['after'] for x in json.loads((P/'editorial_changes.json').read_text())}
postfreeze={x['before']:x['after'] for x in json.loads((P/'postfreeze_edits.json').read_text())}
live={x['before']:x['after'] for x in json.loads((P/'live_integration_edits.json').read_text())}
latency={x["before"]:x["after"] for x in json.loads((P/"receipt_latency_edits.json").read_text())}
cooja={x['before']:x['after'] for x in json.loads((P/'cooja_integration_edits.json').read_text())}
remediation={x['before']:x['after'] for x in json.loads((P/'remediation_edits.json').read_text())['replacements']} if remediated else {}
def after_editorial(s):
 u=postfreeze.get(s,s);v=live.get(u,u);w=latency.get(v,v);result=transform_text(cooja.get(w,w));return '\n'.join(remediation.get(part,part) for part in result.split('\n'))
def final_text(s):
 return after_editorial(editorial.get(transform(s),transform(s)))
def citations(s):
 out=[]
 for x in re.findall(r'\[([\d, ]+)\]',s):
  nums=[int(v) for v in x.split(',')]
  if all(1<=v<=21 for v in nums):out.append(nums)
 return out
for e in revision['paragraphs']:
 replacement=' '.join(e['replacement']);assert citations(e['match'])==citations(replacement),'Citation changed in prose edit'
 before=set(re.findall(r'\d+(?:\.\d+)?',e['match']));after=set(re.findall(r'\d+(?:\.\d+)?',replacement))
 removed=before-after
 assert not removed or (removed=={'30'} and replacement.startswith('Thirty fixtures')),'Numeric evidence removed: '+str(removed)
 for para in e['replacement']:assert final_text(para) in text,'Edited paragraph missing'
for changed in editorial.values():assert after_editorial(changed) in text,'Final editorial paragraph missing'
for key in ['shibambhardwaj@gmail.com','0009-0005-4554-7397','abhishek.pandey2@bennett.edu.in','0000-0003-3799-9754','abhinandan282@bit.ac.in','No funding was received for this research.','The authors declare no competing interests.']:assert key in text
assert not any(x in text for x in ['Jitendra Nath Srivastava','Akash Sanghi','Invertis University'])
abstract=d.paragraphs[next(i for i,p in enumerate(d.paragraphs) if p.text=='Abstract')+1].text;assert len(abstract.split())<=250
with zipfile.ZipFile(P/'Array_Manuscript.docx') as z:
 embedded={hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
 exports={hashlib.sha256(p.read_bytes()).hexdigest() for folder in (['figures','review_figures','postfreeze_figures','cooja_figures']+(['remediation_figures'] if remediated else [])) for p in (P/folder).glob('fig*.png')}
 assert embedded==exports and len(embedded)==expected_figures
result=dict(status='passed',prose_paragraphs_edited=len(revision['paragraphs']),native_equations_unchanged=10,frozen_tables_unchanged=5,transfer_table_reconciled=True,figures_embedded=expected_figures,tables_inline=expected_tables,references=25,abstract_words=len(abstract.split()),citations_preserved=True,numeric_tokens_preserved_except_thirty_word_conversion=True,all_exhibits_before_references=True,images_match_verified_exports=True,physical_validation=False)
(P/'manuscript_audit.json').write_text(json.dumps(result,indent=2)+'\n');print(json.dumps(result,indent=2))
