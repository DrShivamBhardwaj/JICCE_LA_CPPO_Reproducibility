"""Embed only reconciled post-review diagnostics, with explicit scope."""
import json
from pathlib import Path
from docx.shared import Cm,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def apply(doc,root):
 a=json.loads((root/'outputs/review_extensions/analysis.json').read_text());assert a['long_runs']==18 and a['sensitivity_runs']==33
 anchor=next(p for p in doc.paragraphs if p.text=='6. Discussion')
 def para(text,style=None):
  p=anchor.insert_paragraph_before(text,style);p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(6);p.paragraph_format.widow_control=True;return p
 para('5.5 Post review stress and sensitivity diagnostics','Heading 2')
 para('Table 7 and Fig. 9 report a separate exploratory extension on fixtures 84001–84003, using attention checkpoints 41001–41003. Active proofs, censored passive and frozen prior each run for 60 and 300 s, giving 18 runs. Packets arrive every 0.2 s from 0.4 s until near the end of the run: 290 and 1,490 offered packets, respectively. The topology, blackhole fraction, beacon schedule and 0.5 J routing-energy scale follow the frozen benchmark. These diagnostics were designed after review and do not enlarge the original confirmatory test family.')
 para('The extension exposes a limit rather than steady operation. At the assumed listening current, 0.5 J corresponds to roughly 8.87 s. Routing excludes an energy-depleted node, but the radio bridge continues telemetry and PHY-state integration. Delivery consequently plateaus while offered traffic and source-service debt continue to grow. Accumulated radio energy after depletion is bookkeeping for an energized simulation, not energy physically supplied by the original battery. Retained-watch counts also persist because closed records are not reclaimed.')
 p=para('Table 7. Long horizon stress diagnostics.');p.paragraph_format.keep_with_next=True
 table=doc.add_table(rows=1,cols=5);table.autofit=False;anchor._p.addprevious(table._tbl)
 for col,width in zip(table.columns,[4.8,1.8,2.3,3.4,3.2]):col.width=Cm(width)
 headers=['Horizon / arm','Packets','Delivery (%)','Critical failure (%)','Max. source debt']
 for c,v in zip(table.rows[0].cells,headers):c.text=v
 labels={'active_proof':'Active proofs','censored_passive':'Censored passive','frozen_prior':'Frozen prior'}
 for r in a['long_groups']:
  vals=[f"{r['duration_s']} s / {labels[r['arm']]}",str(r['packets_per_run']),f"{r['delivery_mean_percent']:.2f}",f"{r['critical_failure_mean_percent']:.2f}",f"{r['source_debt_max']:.2f}"]
  for c,v in zip(table.add_row().cells,vals):c.text=v
 # Journal horizontal rules; no split rows; keep the compact table together.
 for i,row in enumerate(table.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for cell,width in zip(row.cells,[4.8,1.8,2.3,3.4,3.2]):
   cell.width=Cm(width)
   for p in cell.paragraphs:
    p.paragraph_format.line_spacing=1.05;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
    for run in p.runs:run.font.size=Pt(9);run.bold=i==0
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','bottom','insideH']:
  e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'888888');borders.append(e)
 table._tbl.tblPr.append(borders)
 para('Delivery and critical-failure entries are means across three fixtures; debt is the maximum final delay debt across sources and fixtures. Zero handle-exhaustion events were recorded. Pending-map samples are taken once per second and cannot rule out transient MAC queue drops. This small extension supplies no bounded-queue or stability guarantee.')
 def figure(name,caption):
  p=para('');p.add_run().add_picture(str(root/'publication/array/review_figures'/f'{name}.png'),width=Cm(16));p.paragraph_format.keep_with_next=True
  p=para(caption)
  for run in p.runs:run.font.size=Pt(10)
 figure('fig9_long_horizon','Fig. 9. Executed 300 s stress trajectories.')
 effects=a['sensitivity_effects'];span=[min(x['range'][0] for x in effects),max(x['range'][1] for x in effects)]
 para(f'Figure 10 examines Eq. (1) on three further fixtures, 85001–85003, paired with checkpoints 41001–41003. Each coefficient is changed alone by −20% or +20%; the unchanged setting and ten perturbations give 33 four-second blackhole runs. All 30 perturbed fixture-level delivery contrasts are exactly zero in these runs. This limited result does not establish insensitivity in denser candidate sets or outside the tested range. These are local sensitivity observations, not a tuning procedure or a confidence interval. No checkpoint is retrained and no setting is selected from these results.')
 figure('fig10_weight_sensitivity','Fig. 10. Executed coefficient sensitivity contrasts.')
 para('The full actor–critic contains 63,368 parameters, while the inference actor requires 44,161 parameters in the audited implementation: 176,644 FP32 weight bytes before runtime buffers. This arithmetic corrects treating the complete training network as an unavoidable embedded inference footprint. It does not establish an MCU deployment. Duty-cycled MAC performance, drift tolerance, target inference and measured energy require separate validation.')
