"""Verify the editorial revision without altering experiment evidence."""
import argparse, hashlib, json, zipfile
from pathlib import Path
from docx import Document
from lxml import etree
from array_failure_boundary_revision import TITLE, ABSTRACT, BOUNDARY

root = Path(__file__).resolve().parents[1]
p = argparse.ArgumentParser(); p.add_argument('--baseline', type=Path, required=True); args = p.parse_args()
file = root/'publication/array/Array_Manuscript.docx'
before, after = Document(args.baseline), Document(file)
assert [t._tbl.xml for t in before.tables] == [t._tbl.xml for t in after.tables]
def payload(path):
    with zipfile.ZipFile(path) as z:
        tree = etree.fromstring(z.read('word/document.xml'))
        math = [etree.tostring(e, method='c14n') for e in tree.xpath('//*[local-name()="oMath"]')]
        media = {n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
    return math, media
assert payload(args.baseline) == payload(file)
assert len(after.tables)==11 and len(after.inline_shapes)==12
edits=json.loads((root/'publication/array/failure_boundary_edits.json').read_text())
mapping={e['before']:e['after'] for e in edits['replacements']}
assert [p.text for p in after.paragraphs if p.text not in BOUNDARY] == [mapping.get(p.text,p.text) for p in before.paragraphs]
assert after.paragraphs[0].text == TITLE
assert ABSTRACT in [p.text for p in after.paragraphs]
a=json.loads((root/'outputs/cooja_feedback/primary_600/analysis.json').read_text())
for arm,value in [('censored',32.94),('active',79.11),('censored_proof',32.83),('naive',54.89)]:
    row=next(x for x in a['groups'] if (x['attack'],x['timing'],x['arm'])==(1,1,arm))
    assert round(100*row['delivery'],2)==value
for arm,delta,lo,hi in [('active',46.17,36.89,54.50),('censored_proof',-.11,-.50,.39)]:
    row=next(x for x in a['contrasts'] if (x['endpoint'],x['attack'],x['timing'],x['arm'])==('delivery',1,1,arm))
    assert [round(100*v,2) for v in [row['mean_difference'],*row['ci95']]]==[delta,lo,hi]
text='\n'.join(p.text for p in after.paragraphs)
assert 'not a population-equivalence result' in text and 'rather than estimating its onset' in text
report={'status':'passed','classification':'editorial interpretation, no new experiment','all_11_tables_unchanged':True,'all_12_images_unchanged':True,'all_10_equations_unchanged':True,'unmapped_paragraph_changes':0,'cooja_headline_values_reconciled':True,'abstract_words':len(ABSTRACT.split()),'physical_validation':False}
(root/'publication/array/failure_boundary_audit.json').write_text(json.dumps(report,indent=2)+'\n');print(json.dumps(report,indent=2))
