"""Integrate preserved live-MCU evidence without changing frozen statistics."""
from pathlib import Path
from copy import deepcopy
import json
from docx.shared import Pt,Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def apply(doc,root):
 root=Path(root);changes=[]
 def replace(prefix,fn):
  p=next(p for p in doc.paragraphs if p.text.startswith(prefix));old=p.text;p.text=fn(old);changes.append(dict(before=old,after=p.text))
 replace("Renode's deterministic range medium",lambda s:s.replace('embedded policy/crypto execution or calibrated current consumption.','calibrated current consumption. Section 5.8 adds a separate live MCU policy/crypto pilot; this earlier host-driven fixture is unchanged.'))
 replace('The full attention actor',lambda s:s.replace('The present MCU firmware handles','The host-controlled firmware in this fixture handles'))
 replace('A four-second horizon, a central sink',lambda s:s.replace('integrated firmware costs remain unmeasured','physical integrated-firmware costs remain unmeasured'))
 replace('The image omits radio buffers',lambda s:s+' Section 5.8 subsequently evaluates a combined radio and policy/crypto image under a narrower four-node protocol.')
 replace('The inference actor contains 44,161',lambda s:s.replace('Duty-cycled MAC performance, drift tolerance, integrated target inference and measured energy require separate validation.','Section 5.8 tests integrated target inference in emulation. Duty-cycled MAC performance, physical drift tolerance and measured energy still require separate validation.'))
 replace('ACK-gated censored passive feedback passed the prespecified',lambda s:s.replace('while inference and cryptography remain on the host.','while inference and cryptography run on the host in that earlier fixture.').replace('integrated on-board policy/crypto execution and measured energy remain unvalidated.','physical on-board policy/crypto execution and measured energy remain unvalidated.')+' A four-node emulated live-path pilot now performs routing and authentication inside the MCU, including longer private-key runs. Successful delivery does not establish real-time service: receipt delays exceed the two tighter class budgets.')
 replace('A useful next step is to decide' if any(p.text.startswith('A useful next step is to decide') for p in doc.paragraphs) else 'Section 5.6 separates proof generation',lambda s:s.replace('absent from the isolated harness.','not established by isolated or small-scale integrated emulation.'))
 abstract=doc.paragraphs[next(i for i,p in enumerate(doc.paragraphs) if p.text=='Abstract')+1];old=abstract.text
 abstract.text=old.replace('Compiled UART-to-radio firmware','Earlier UART-to-radio firmware').replace('An isolated Cortex-M3 harness reproduces actor decisions and executes HMAC-SHA256 and Ed25519.','A four-node MCU emulation pilot executes live routing and authentication but misses the two tighter receipt deadlines.')
 changes.append(dict(before=old,after=abstract.text));assert len(abstract.text.split())<=250
 anchor=next(p for p in doc.paragraphs if p.text=='6. Discussion');body=next(p for p in doc.paragraphs if p.text.startswith('The simulated network has'))._p.pPr
 def para(text,style=None):
  p=doc.add_paragraph(text,style=style)
  if style is None:
   if p._p.pPr is not None:p._p.remove(p._p.pPr)
   p._p.insert(0,deepcopy(body))
  anchor._p.addprevious(p._p)
  if style or text.startswith('Table '):p.paragraph_format.keep_with_next=True
  return p
 runs=[]
 for folder in ['outputs/cc2538_live/chain','outputs/cc2538_live/choice','outputs/cc2538_v2/chain_100','outputs/cc2538_v2/choice_100']:
  d=json.loads((root/folder/'verification.json').read_text());assert d['status']=='passed';runs.append(d)
 para('5.8 Live MCU routing and authentication','Heading 2')
 para('A separate four-node CC2538 pilot moves observation construction, frozen attention inference, hop HMAC, replay checks and signed-receipt processing into the compiled MCU image. The host triggers traffic and authenticated beacon exchange and records UART output; it does not choose ordinary next hops or seal ordinary data and receipts. Node-local cache, link estimates and service debts reconstruct the 93-element observation. The original three-hop chain and a geometry with two feasible source choices each offer ten packets. Table 10 separates these public-fixture runs from two longer, separately provisioned runs. These functional fixtures are not independent statistical replicates of the original benchmark.')
 para('Every MCU decision matches the frozen Python actor, and an independent feature reconstruction differs by less than 1e-4. Captured ordinary hop tags, source authenticators and sink signatures are checked independently. Exact replay, a corrupt hop tag, a corrupt source tag inside a valid hop envelope, and a forged sink signature inside a valid hop envelope are rejected. The longer version additionally rejects an authenticated stale boot session and repeated provisioning. It receives fresh random pairwise keys and a sink signing seed over the trusted emulator-control UART; a locked, persisted host allocator reserves new 64-bit boot sessions before startup. Private audit fixtures remain outside the public repository. These are emulator provisioning controls, not protected MCU storage or persistent device firmware.')
 para('Table 10. Live MCU pilot results.')
 headers=['Topology','Offered','Delivered / receipts','MCU decisions','Receipt delay (ms)','Deadline misses']
 rows=[]
 for d in runs:
  delays=d['emulated_receipt_delays'];miss=sum(x['receipt_delay_ms']>[90,350,1500][x['traffic']] for x in delays)
  rows.append([('Chain' if d['topology']=='chain' else 'Choice')+(' private' if d['expected_packets']==100 else ' public'),str(d['expected_packets']),f"{d['unique_sink_deliveries']} / {d['unique_source_verified_receipts']}",str(d['mcu_policy_decisions']),f"{min(x['receipt_delay_ms'] for x in delays)}–{max(x['receipt_delay_ms'] for x in delays)}",str(miss)])
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
 for row in [headers]+rows:
  cells=t.rows[0].cells if row is headers else t.add_row().cells
  for c,v,w in zip(cells,row,[2.7,1.4,2.5,2.2,3.0,2.1]):
   c.width=Cm(w);c.text=v
   for p in c.paragraphs:
    p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1
    for run in p.runs:run.font.name='Times New Roman';run.font.size=Pt(8.5);run.bold=row is headers
 for i,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if i==0:pr.append(OxmlElement('w:tblHeader'))
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','bottom','insideH','left','right','insideV']:
  e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single' if edge in ['top','bottom','insideH'] else 'nil');e.set(qn('w:sz'),'4');borders.append(e)
 t._tbl.tblPr.append(borders);anchor._p.addprevious(t._tbl)
 para('The first kernel evaluated masked candidates unnecessarily and allowed relay cache expiry; its failed development capture is retained. Skipping only masked computations restores the declared freshness behavior without changing weights or extending timing windows. The O2-compiled version uses a rolling source-identity replay window to support longer sequential streams; it does not establish deadline compliance. The original runs take 1119–1128 ms in the chain and 915–922 ms in the choice geometry; seven of ten receipts in each miss their 90 or 350 ms class budget. Table 10 reports the longer-run delays without treating delivery as real-time success. All times come from the Renode SysTick model, not a physical board measurement.')
 para('The pilot fixes source 0, sink 3 and four provisioned nodes, with sixteen pending records. Beacon energy and queue inputs are nominal, not measured telemetry; beacon scheduling is test-triggered. Longer streams test state reuse and replay-window progression rather than network-size scaling. Worst-case stack, overload, interference, autonomous rendezvous, real board startup, reboot-safe MCU storage and measured RF/energy remain outside this validation. The private-key version requires locally retained private fixtures for retrospective hop-tag audit; public-fixture captures remain fully inspectable. None of these runs supports a neural-superiority or deployment-readiness claim.')
 (root/'publication/array/live_integration_edits.json').write_text(json.dumps(changes,indent=2,ensure_ascii=False)+'\n')
 return dict(tables=10,figures=11,references=24,abstract_words=len(abstract.text.split()))
