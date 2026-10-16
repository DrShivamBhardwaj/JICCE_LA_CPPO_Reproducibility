"""Focused JICCE review manuscript with native math and trace-linked exhibits."""
from pathlib import Path
from copy import deepcopy
import json,re,zipfile
from collections import OrderedDict
import numpy as np
from lxml import etree
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK,WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'publication/redesigned';OUT.mkdir(parents=True,exist_ok=True)
A=json.loads((ROOT/'outputs/ns3_censored/paired_analysis.json').read_text());ROWS=json.loads((ROOT/'outputs/ns3_censored/summary.json').read_text());AUD=json.loads((ROOT/'outputs/ns3_censored/implementation_audit.json').read_text())
REFS={
'rpl':('T. Winter et al., “RPL: IPv6 Routing Protocol for Low-Power and Lossy Networks,” IETF, RFC 6550, 2012. DOI: 10.17487/RFC6550.','https://www.rfc-editor.org/rfc/rfc6550','primary standard'),
'etx':('D. S. J. De Couto, D. Aguayo, J. Bicket, and R. Morris, “A high-throughput path metric for multi-hop wireless routing,” in Proceedings of MobiCom, pp. 134–146, 2003. DOI: 10.1145/938985.939000.','https://doi.org/10.1145/938985.939000','established reference retained from verified source manuscript'),
'ctp':('O. Gnawali, R. Fonseca, K. Jamieson, D. Moss, and P. Levis, “Collection tree protocol,” in Proceedings of SenSys, pp. 1–14, 2009. DOI: 10.1145/1644038.1644040.','https://collaborate.princeton.edu/en/publications/collection-tree-protocol/','author institution abstract and metadata'),
'watchdog':('S. Marti, T. J. Giuli, K. Lai, and M. Baker, “Mitigating routing misbehavior in mobile ad hoc networks,” in Proceedings of MobiCom, pp. 255–265, 2000. DOI: 10.1145/345910.345955.','https://sites.cs.ucsb.edu/~ebelding/courses/290I/w03/papers/Marti_Mitigating.pdf','primary full text'),
'ack':('K. Liu, J. Deng, P. K. Varshney, and K. Balakrishnan, “An acknowledgment-based approach for the detection of routing misbehavior in MANETs,” IEEE Transactions on Mobile Computing, vol. 6, no. 5, pp. 536–550, 2007. DOI: 10.1109/TMC.2007.1036.','https://doi.org/10.1109/TMC.2007.1036','indexed primary abstract; full text unavailable during review'),
'edge':('A. Saleh, P. Joshi, R. S. Rathore, and S. S. Sengar, “Trust-aware routing mechanism through an edge node for IoT-enabled sensor networks,” Sensors, vol. 22, no. 20, Art. no. 7820, 2022. DOI: 10.3390/s22207820.','https://www.mdpi.com/1424-8220/22/20/7820','SciSpace discovery; publisher abstract and indexed metadata'),
'path':('N. Kumar and S. Sharma, “PathGuard: Trustworthy routing for sustainable and secure IoT-WSN networks,” Wireless Personal Communications, vol. 136, pp. 469–487, 2024. DOI: 10.1007/s11277-024-11289-8.','https://link.springer.com/article/10.1007/s11277-024-11289-8','SciSpace discovery; publisher abstract only'),
'smart':('P. Goswami, T. Khan, V. Pathak, and A. Alabdultif, “Machine learning based dynamic trust estimation framework for securing wireless sensor networks,” Scientific Reports, vol. 15, Art. no. 35821, 2025. DOI: 10.1038/s41598-025-19768-z.','https://www.nature.com/articles/s41598-025-19768-z.pdf','primary full text'),
'frdr':('S. Tao, H. Tang, J. Wang, and B. Li, “A failure risk-aware multi-hop routing protocol in LPWANs using deep Q-network,” Sensors, vol. 25, no. 14, Art. no. 4416, 2025. DOI: 10.3390/s25144416.','https://www.mdpi.com/1424-8220/25/14/4416','publisher abstract and metadata'),
'rlfs':('A. A. Lata and M. Kang, “RLFS-OR: Reinforcement learning-based forwarder selection for opportunistic routing in wireless sensor networks,” Electronics, vol. 15, no. 5, Art. no. 910, 2026. DOI: 10.3390/electronics15050910.','https://www.mdpi.com/2079-9292/15/5/910','publisher indexed abstract; details not independently replicated'),
'tgrl':('R. Zheng, D. Yamkhin, J. Wei, H. Liu, and Z. Khurelbaatar, “Trust-gated reinforcement learning-based secure routing for wireless sensor networks,” Journal of Computational and Cognitive Engineering, online first, 2026. DOI: 10.47852/bonviewJCCE62029141.','https://ojs.bonviewpress.com/index.php/JCCE/article/view/9141','primary publisher abstract'),
'drlsecure':('R. Gupta, S. K. Gupta, S. B. Singh, R. Kumar, and R. Muthukkumar, “Adaptive deep reinforcement learning–based secure routing for wireless sensor networks,” Scientific Reports, 2026. DOI: 10.1038/s41598-026-62519-x.','https://www.nature.com/articles/s41598-026-62519-x','publisher indexed abstract and metadata'),
'ppo':('J. Schulman, F. Wolski, P. Dhariwal, A. Radford, and O. Klimov, “Proximal policy optimization algorithms,” arXiv:1707.06347, 2017. DOI: 10.48550/arXiv.1707.06347.','https://arxiv.org/abs/1707.06347','primary abstract and manuscript'),
'gae':('J. Schulman, P. Moritz, S. Levine, M. I. Jordan, and P. Abbeel, “High-dimensional continuous control using generalized advantage estimation,” in Proceedings of ICLR, 2016. Available: https://arxiv.org/abs/1506.02438.','https://arxiv.org/abs/1506.02438','primary abstract and manuscript'),
'sets':('M. Zaheer et al., “Deep sets,” in Advances in Neural Information Processing Systems, vol. 30, pp. 3391–3401, 2017. Available: https://papers.nips.cc/paper/6931-deep-sets.','https://papers.nips.cc/paper/6931-deep-sets','primary proceedings'),
'neely':('M. J. Neely, Stochastic Network Optimization with Application to Communication and Queueing Systems. Morgan and Claypool, 2010. DOI: 10.2200/S00271ED1V01Y201006CNT007.','https://www.morganclaypoolpublishers.com/catalog_orig/product_info_products_id-134.html','primary publisher metadata'),
'cpo':('J. Achiam, D. Held, A. Tamar, and P. Abbeel, “Constrained policy optimization,” in Proceedings of ICML, vol. 70, pp. 22–31, 2017. Available: https://proceedings.mlr.press/v70/achiam17a.html.','https://proceedings.mlr.press/v70/achiam17a.html','primary proceedings'),
'ns3':('ns-3 Project, “IEEE 802.15.4 low-rate wireless personal area network model,” ns-3 Model Library, release 3.44. Available: https://www.nsnam.org/docs/release/3.44/models/ns-3-model-library.pdf. Accessed October 16, 2026.','https://www.nsnam.org/docs/release/3.44/models/ns-3-model-library.pdf','official versioned model documentation'),
'holm':('S. Holm, “A simple sequentially rejective multiple test procedure,” Scandinavian Journal of Statistics, vol. 6, no. 2, pp. 65–70, 1979. Available: https://www.jstor.org/stable/4615733.','https://www.jstor.org/stable/4615733','primary archive metadata'),
}
order=OrderedDict();md=[];d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
for side in ['top_margin','bottom_margin','left_margin','right_margin']:setattr(sec,side,Cm(2.5))
normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(9.5);normal.font.color.rgb=RGBColor(0,0,0);normal.paragraph_format.line_spacing=2;normal.paragraph_format.space_after=Pt(4)
for name,size in [('Title',15),('Heading 1',11),('Heading 2',10)]:
 st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.keep_with_next=True
 d.styles[name].paragraph_format.space_before=Pt(10);d.styles[name].paragraph_format.space_after=Pt(5)
d.styles['Heading 1'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;d.styles['Heading 2'].font.italic=True
for p in sec.footer.paragraphs:
 p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run();f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');r._r.addnext(f)
def cite(s):
 def f(m):
  key=m.group(1);assert key in REFS,key
  if key not in order:order[key]=len(order)+1
  return '['+str(order[key])+']'
 return re.sub(r'\{\{([a-z0-9]+)\}\}',f,s)
def p(s,style=None):
 s=cite(s);x=d.add_paragraph(s,style);x.paragraph_format.widow_control=True;md.append(s+'\n');return x
def h(s,level=1):md.append('#'*(level+1)+' '+s+'\n');return d.add_paragraph(s,'Heading '+str(level))
M='http://schemas.openxmlformats.org/officeDocument/2006/math'
def mr(s):
 x=etree.Element('{'+M+'}r');etree.SubElement(x,'{'+M+'}t').text=str(s);return x
def sub(a,b):
 x=etree.Element('{'+M+'}sSub');etree.SubElement(x,'{'+M+'}e').append(mr(a));etree.SubElement(x,'{'+M+'}sub').append(mr(b));return x
def sup(a,b):
 x=etree.Element('{'+M+'}sSup');etree.SubElement(x,'{'+M+'}e').append(mr(a));etree.SubElement(x,'{'+M+'}sup').append(mr(b));return x
def frac(a,b):
 x=etree.Element('{'+M+'}f')
 for tag,value in [('num',a),('den',b)]:
  e=etree.SubElement(x,'{'+M+'}'+tag)
  for item in (value if isinstance(value,list) else [value]):e.append(mr(item) if isinstance(item,str) else item)
 return x
def summ(index,items):
 x=etree.Element('{'+M+'}nary');pr=etree.SubElement(x,'{'+M+'}naryPr');etree.SubElement(pr,'{'+M+'}chr').set('{'+M+'}val','∑');etree.SubElement(pr,'{'+M+'}limLoc').set('{'+M+'}val','undOvr');etree.SubElement(pr,'{'+M+'}supHide').set('{'+M+'}val','1');etree.SubElement(x,'{'+M+'}sub').append(mr(index));etree.SubElement(x,'{'+M+'}sup');e=etree.SubElement(x,'{'+M+'}e')
 for item in items:e.append(mr(item) if isinstance(item,str) else item)
 return x
n_eq=0
def eq(items,plain,defs):
 global n_eq;n_eq+=1;x=d.add_paragraph();x.alignment=WD_ALIGN_PARAGRAPH.LEFT;x.paragraph_format.keep_with_next=True;x.paragraph_format.line_spacing=1.6;x.paragraph_format.tab_stops.add_tab_stop(Cm(8),WD_TAB_ALIGNMENT.CENTER);x.paragraph_format.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT);x.add_run('\t');math=etree.SubElement(x._p,'{'+M+'}oMath')
 for v in items:math.append(mr(v) if isinstance(v,str) else v)
 x.add_run('\t('+str(n_eq)+')');md.append('Equation '+str(n_eq)+': '+plain+'\n')
 for text in defs:
  z=p('• '+text);z.paragraph_format.line_spacing=1.15;z.paragraph_format.space_after=Pt(3)
def page():d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
def table(num,title,heads,rows,widths):
 page();z=p('Table '+str(num)+'. '+title);z.paragraph_format.keep_with_next=True;z.paragraph_format.line_spacing=1
 t=d.add_table(rows=1,cols=len(heads));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,x in zip(t.rows[0].cells,heads):c.text=cite(x)
 for row in rows:
  cells=t.add_row().cells
  for c,x in zip(cells,row):c.text=cite(str(x))
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','bottom','insideH','left','right','insideV']:
  e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single' if edge in ['top','bottom'] else 'nil');e.set(qn('w:sz'),'6');e.set(qn('w:color'),'777777');borders.append(e)
 t._tbl.tblPr.append(borders)
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for i,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  for j,c in enumerate(row.cells):
   c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   for para in c.paragraphs:
    para.paragraph_format.line_spacing=1.25;para.paragraph_format.space_before=Pt(4);para.paragraph_format.space_after=Pt(4)
    para.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
    for run in para.runs:run.font.name='Times New Roman';run.font.size=Pt(8.5);run.bold=i==0 or str(c.text).startswith('Censored passive')
 md.append('| '+' | '.join(heads)+' |\n|'+'---|'*len(heads)+'\n'+'\n'.join('| '+' | '.join(map(str,row))+' |' for row in rows)+'\n')
def fig(num,name,caption):
 page();x=d.add_paragraph();x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.line_spacing=1;x.paragraph_format.keep_with_next=True;x.add_run().add_picture(str(OUT/'figures'/f'{name}.png'),width=Cm(16));x=p('Fig. '+str(num)+'. '+caption);x.paragraph_format.line_spacing=1.15;md.append('![Fig. '+str(num)+'](figures/'+name+'.svg)\n')
# Focused narrative. Historical evidence remains an explicitly linked supplement.
p('Receipt Assisted Passive Feedback for Multihop Wireless Sensor Routing','Title')
p('Shivam Bhardwaj, Jitendra Nath Srivastava and Akash Sanghi').alignment=WD_ALIGN_PARAGRAPH.CENTER
p('Faculty of Engineering and Technology, Invertis University, Bareilly, Uttar Pradesh 243123, India').alignment=WD_ALIGN_PARAGRAPH.CENTER
p('Corresponding author: shibambhardwaj@gmail.com\nCoauthors: jitendranathshrivastava@gmail.com; sanghiakash@gmail.com')
p('Running title: Passive Feedback for Sensor Routing')
h('Abstract')
abstract='Routing feedback consumes the channel it monitors and can confuse missing observations with forwarding failure. We evaluate receipt assisted passive feedback for a frozen local attention routing policy in an IEEE 802.15.4 simulator. A development study exposed adverse effects of explicit forwarding proofs. The redesigned ledger was then frozen before 750 runs on 30 new fixtures, using naive passive, active proof, frozen prior and capacity matched policy controls. Under blackholes, delivery increased from 71.28% with active proofs to 75.94%, a paired gain of 4.67 percentage points with a 95% confidence interval of 2.36–6.97 and Holm adjusted p=0.0161. All four prespecified incremental gate checks passed. Neither censoring nor attention established superiority over its matched controls. Energy per delivered byte improved, while total modeled radio energy increased slightly. The results support a bounded replacement of the proof based mechanism under the tested assumptions, with open limits in authentication, distributed state, duty cycling and absolute service compliance.'
assert 140<=len(abstract.split())<=160,len(abstract.split());p(abstract)
p('Index Terms: Feedback censoring, Multihop routing, Passive monitoring, Reproducibility, Wireless sensor networks')
h('I. Introduction')
p('Multihop sensor routing relies on link estimates, path progress and control messages, as illustrated by RPL, expected transmission count and the Collection Tree Protocol {{rpl}}{{etx}}{{ctp}}. Adding security evidence creates another dependency: an extra acknowledgment may increase confidence in one hop while occupying airtime needed by another packet. A mechanism should therefore be judged jointly by what it observes and what it sends.')
p('Passive watchdog monitoring and explicit acknowledgment schemes have addressed forwarding misbehavior for decades {{watchdog}}{{ack}}. Overhearing establishes that a relay attempted transmission; it does not establish reception by the intended next hop. Conversely, silence can result from an unsuccessful first link, contention or missed reception. Treating every silence as a malicious action can change routing even in an attack free network. The present study examines that observation ambiguity without introducing an intrusion classifier.')
p('Recent work links trust to edge assisted routing {{edge}}, optimizer based next hop selection {{path}}, learned trust estimation {{smart}}, failure aware deep Q learning {{frdr}} and opportunistic forwarder selection {{rlfs}}. Trust gated reinforcement learning {{tgrl}} and secure deep reinforcement learning {{drlsecure}} further demonstrate current interest in the interaction between security and routing. Table 1 maps their stated mechanisms to this study. Their reported performance values are not used as directly comparable baselines because the channels, workloads and attack models differ.')
p('Our initial architecture studies did not establish joint superiority over a capacity matched multilayer perceptron. A subsequent 600-run radio telemetry study also found no corrected architecture advantage and exposed a cost of the active proof mechanism. These negative stages motivated one feedback redesign; they remain available with their raw outputs. The new experiment asks a narrower question: can a receipt assisted passive mechanism replace active proofs without materially degrading a frozen prior control on a new fixture set?')
h('A. Contribution and research questions',2)
p('The contribution is an auditable feedback-mechanism comparison: first-hop acknowledgment gates evidence, a record retains the strongest observation rather than every duplicate, and missing observations remain censored. The implementation connects these rules to actual packet reception callbacks. It does not claim to invent watchdog monitoring, censoring, attention or reinforcement learning. Figure 1 separates radio events from simulator-held state, and Fig. 2 shows the event logic.')
p('Three questions organize the evaluation. RQ1 asks whether the replacement improves blackhole delivery relative to active proofs under the frozen multiple-testing rule. RQ2 asks whether it remains within prespecified delivery, critical-service and energy-efficiency margins against frozen-prior attention. RQ3 asks whether the naive passive and capacity matched controls support attribution to censoring or architecture. These questions distinguish a useful replacement from a unique algorithmic advantage.')
h('II. System model and methods')
h('A. Radio traffic and threat model',2)
p('The network contains 100 stationary sensors and one sink at (75, 75) m in a 150 × 150 m field. The ns-3.44 low-rate wireless personal area network model supplies unslotted medium access, acknowledgments and packet-level reception {{ns3}}. Log-distance loss uses exponent 3 and reference loss 40 dB; the weak-link condition uses 54 dB. A separate 50 m range limit is retained. This is a radio simulation, not a measurement of an operational sensor deployment.')
p('Each run lasts 4 s and offers 60 source packets after a 0.4 s warm-up. Source intervals are 20 ms, or 2 ms for the congested burst. Traffic classes occur with probabilities 0.35, 0.45 and 0.20 and deadlines of 90, 350 and 1,500 ms. A 64-byte application frame contains 12 identifier bytes and 52 payload bytes; lower-layer overhead is additional. Periodic application-level telemetry broadcasts are termed beacons here; they are not IEEE superframe beacons. They occur every 0.5 s with a 3 ms node offset. Table 2 records the common settings.')
p('In blackhole settings, 25% of sensor identities are selected independently for each fixture. A faulty relay accepts a packet but does not invoke its forwarding decision. It continues to produce the modeled beacon and receipt behavior; forged control values, collusion and compromised sink behavior are outside the threat model. Fault labels control the simulator and evaluation only. The policy receives forwarding observations rather than node attack labels.')
h('B. State construction and frozen routing policy',2)
p('A received beacon admits an observer-specific cache entry with energy fraction, pending-request pressure, timestamp and link-quality indication. An entry expires after 0.75 s. Candidate construction excludes the current node, previously visited nodes, exhausted nodes and nodes beyond 50 m. At most eight candidates are retained using the fixed preselection score')
eq([sub('S','uv'),' = 1.25 ',sub('g','uv'),' + 0.75 ',sub('q','uv'),' + 0.25 ',sub('e','v'),' − 0.55 ',sub('R','uv'),' − 0.20 ',sub('Q','v')], 'S_uv = 1.25 g_uv + 0.75 q_uv + 0.25 e_v - 0.55 R_uv - 0.20 Q_v', ['g is sinkward progress divided by 50 m; q is estimated MAC success probability; both are dimensionless.','e is remaining modeled energy divided by 0.5 J, Q is pending-request pressure normalized by 32, and R is the evidence-risk upper proxy; all are dimensionless.'])
p('The complete observation has 13 context values and eight ten-value candidate rows. Context includes local energy and queue state, normalized sink distance, remaining deadline, class indicators and six source-service debts. Each candidate row contains a sink indicator, energy fraction, estimated MAC success, received link quality, queue pressure, sinkward progress, evidence risk and uncertainty, expected retries, and normalized link delay. The training RSSI proxy is replaced by received link-quality indication during radio transfer; this feature shift is retained rather than learned away.')
p('Separate shared encoders transform context x and candidate h into 96-dimensional vectors. Query-to-candidate scoring and a shared bias produce the logit')
eq([sub('ℓ','uv'),' = ',frac([sub('q','θ'),'(',sub('x','u'),')ᵀ ',sub('k','θ'),'(',sub('h','uv'),')'],'√96'),' + ',sub('b','θ'),'(huv)'], 'ell_uv = q_theta(x_u)^T k_theta(h_uv)/sqrt(96) + b_theta(h_uv)', ['ℓ is a dimensionless action logit; x and h are the context and candidate feature vectors.','qθ and kθ are learned query/key transforms and bθ is a learned candidate bias; θ is fixed throughout transfer.'])
p('Masked argmax selects the relay. Candidate encoding is shared and the critic uses masked mean pooling, following the set-function motivation in {{sets}}. Attention has 63,368 parameters. The capacity matched perceptron has 63,554 parameters and a 203-unit hidden width. Candidate order is score ranked, and equal scores use the configured identity tie break. Consequently, the implementation supports shared candidate scoring but does not prove invariance of the entire routing pipeline to every node relabeling.')
p('Both checkpoint families were trained on seeds 41001–41030 for 400 episodes per seed using eight updates of 50 episodes and three full-batch PPO epochs per update. Learning rate is 0.0003, clipping 0.2, discount 0.99, generalized-advantage factor 0.95 and entropy coefficient 0.01 {{ppo}}{{gae}}. Six cost channels penalize class-specific deadline violation and loss. Lagrange multipliers update at rate 0.03 with clipping to [0, 50]. This finite-budget Lagrangian training is not the algorithmic guarantee of constrained policy optimization {{cpo}}. No checkpoint or threshold is retuned on the new radio fixtures.')
h('C. Receipt assisted censored evidence',2)
p('Each observer u stores a record for selected relay v and packet p. The record contains its expected first-link success, acknowledgment status, expiry, strongest seen evidence and credited strength. A successfully received overheard data frame supplies strength 0.5. A reverse receipt originating from sink delivery supplies strength 1. The maximum rule prevents repeated receptions from accumulating arbitrary strength:')
eq([sub('s','uvp'),' ← max{',sub('s','uvp'),', 0.5 ',sub('I','overheard'),', ',sub('I','receipt'),'}'], 's_uvp <- max(s_uvp, 0.5 I_overheard, I_receipt)', ['s is the strongest evidence strength, dimensionless and bounded by 1.','I_overheard and I_receipt indicate matched received forwarding frames and reverse receipts, respectively; they are dimensionless binary observations.'])
p('Credit is eligible only after the observer obtains a successful MAC acknowledgment for its own data request and the new strength exceeds the amount previously credited. An early observation waits for that acknowledgment. The implemented positive update is')
eq([sub('α','uv'),' ← ',sub('α','uv'),' + clip(',sub('q','uv'), ' max{0, ',sub('s','uvp'),' − ',sub('c','uvp'),'}, 0.05, 1),   ',sub('c','uvp'),' ← ',sub('s','uvp')], 'alpha_uv <- alpha_uv + clip(q_uv max(0,s_uvp-c_uvp),0.05,1); c_uvp <- s_uvp', ['α is a dimensionless effective success count, initialized to 1; c is the credited record strength, initialized to 0.','q is the dimensionless link-success estimate at registration. The clipping floor applies to each positive increment, not to their sum.'])
p('The effective failure count β stays at its prior of 1 in the censored arm. A failed first hop closes the record without positive credit. An acknowledged record with no observed forwarding evidence expires after 150 ms without a failure increment. A late receipt may still upgrade an acknowledged censored record. Near q<0.1, separate half credits interact with the clipping floor and can exceed one full q-weighted increment. The audited decision histories have minimum q=0.6923, so that edge case is outside the observed runs; a calibrated posterior interpretation is nevertheless inappropriate.')
p('The existing risk coordinates are retained as evidence preferences. With the rate-anomaly term inactive, the mean and upper proxy are')
eq([sub('ρ','uv'),' = 0.75(1 − ',frac([sub('α','uv')],[sub('α','uv'),' + ',sub('β','uv')]),'),   ',sub('R','uv'),' = clip(',sub('ρ','uv'),' + 1.645 × 0.75 ',sub('σ','uv'),', 0, 1)'], 'rho_uv = 0.75(1-alpha/(alpha+beta)); R_uv = clip(rho_uv + 1.645*0.75*sigma_uv,0,1)', ['ρ and R are dimensionless risk coordinates.','σ is the beta-form standard deviation sqrt[αβ / ((α+β)²(α+β+1))], used as a dimensionless uncertainty proxy rather than a calibrated maliciousness interval.'])
p('The naive passive control receives identical observations and deduplicates them, but penalizes missing acknowledged forwarding at expiry. The active control uses separate proof packets and its earlier timeout convention; it can credit a proof without applying the passive first-hop ACK gate. Frozen prior preserves the same initial risk coordinates but disables updates. Capacity matched MLP uses the censored ledger. Table 3 makes these differences explicit. Active versus censored therefore compares complete feedback mechanisms, including semantics and traffic, rather than isolating airtime alone.')
h('D. Source service state and simulator boundaries',2)
p('Per-class service debt follows the virtual-queue form motivated by long-run constraint tracking {{neely}}:')
eq([sub('Z','c'),' ← max{0, ',sub('Z','c'),' + ',sub('I','violation'),' − ',sub('b','c'),'}'], 'Z_c <- max(0,Z_c + I_violation - b_c)', ['Zc is a dimensionless source-service debt and bc is the allowed violation fraction for class c.','I_violation is 1 for a missing or late delivered packet and 0 otherwise; loss has a separate analogous channel.'])
p('In this radio transfer, only the source updates those debts after receipt or a 1.5 s expiry. Intermediate nodes do not update a complete service debt for their transit packets. A shared per-packet visited set prevents loops, and a simulator reverse-path map routes receipts. Their state dissemination is not serialized. Beacon energy and queue values likewise use a simulator metadata sidecar keyed to the transmitted identifier; an actual radio receive callback gates availability, but the values are not decoded from wire fields. Figure 1 marks this boundary explicitly. It is a controlled transfer of routing decisions, not an embedded distributed protocol implementation.')
p('Algorithm 1 specifies the tested observer logic. Its event sequence appears in Fig. 2. Physical callback controls check an actual two-hop forwarding observation, silence at an out-of-range observer and zero proof frames in the passive configuration. These checks establish callback behavior; they do not establish authentication or intrusion-detection accuracy.')
h('III. Experimental design')
h('A. Development separation and untouched fixtures',2)
p('Figure 3 accounts for the evidence stages using a PRISMA-inspired staged layout; it is experimental accounting, not a systematic-review selection diagram. The earlier 76001–76030 telemetry fixtures informed the redesign. Protocol, source and the decision rule were published at commit 4027f071322cb5682f39329f656d76df3b68ff62 before evaluation of fixtures 77001–77030. Technical pilot 76901 is excluded from inference. Each new fixture pairs the unchanged attention/MLP checkpoints from 41001 plus its fixture offset.')
p('Five feedback/policy arms cross five settings and 30 fixtures: 750 runs and 45,000 source packets. All runs are retained. Each fixture shares topology, offered packets, traffic classes and blackhole identities across arms. Protocol-induced event schedules can diverge, so identical seeds do not imply identical per-event radio draws. The unit of inference is the complete paired fixture and checkpoint, not individual packets or repeated training of the same network trace. Geometry remains drawn from one stationary placement family with a central sink.')
h('B. Endpoints and radio accounting',2)
p('The primary routing summaries count sink reception against all offered source packets. Delivery fraction is')
eq([ 'D = ',frac([sub('N','delivered')],'60')], 'D = N_delivered / 60', ['D is dimensionless; Ndelivered counts distinct packet identifiers received by the sink during the 4 s horizon.'])
p('For critical class 0, non-delivery counts as a violation, as does delivery beyond 90 ms:')
eq(['V = ',frac([summ('p ∈ C₀',['I[p missing or ',sub('delay','p'),' > 0.09 s]'])],['|',sub('C','0'),'|'])], 'V = sum_{p in C0} I[p missing or delay_p > .09s] / |C0|', ['V is a dimensionless violation fraction and C0 is the set of offered critical-class packets.','delayp is sink reception time minus source generation time, in seconds.'])
p('Custom PHY-state integration charges every sensor’s transmit, receive/listen and off intervals, including control traffic and retries. The sink is recorded separately and excluded from the sensor numerator. The assumed 3 V current model gives')
eq(['E = 3 ',summ('u ∈ sensors',['(0.0174 ',sub('T','u,TX'),' + 0.0188 ',sub('T','u,RX/listen'),' + 0.000020 ',sub('T','u,off'),')'])], 'E = 3 sum_u(0.0174 T_u,TX + 0.0188 T_u,RXlisten + .000020 T_u,off)', ['E is modeled sensor radio energy in joules; residence times T are in seconds and coefficients are amperes.','The supply is 3 V. Currents are explicit assumptions, not calibrated measurements; CPU, sensing and cryptography are excluded.'])
p('Energy efficiency divides this common-horizon expenditure by delivered application payload bytes:')
eq(['η = ',frac('E',['52 ',sub('N','delivered')])], 'eta = E/(52 N_delivered)', ['η is radio joules per delivered payload byte; 52 is payload bytes per packet.','A lower η can result from a larger delivered-byte denominator even if total E does not decrease.'])
h('C. Statistical analysis and frozen gate',2)
p('For each endpoint and comparator, the paired difference is computed once per fixture and summarized across 30 fixtures. Two-sided 95% Student t intervals are unadjusted descriptive intervals. The prespecified 60-test Holm family combines five settings, four comparators and three endpoints {{holm}}. Mean delivery and violation differences are reported in percentage points; energy-efficiency differences use J/byte. The family is not reduced after inspection of results.')
p('All four frozen gate checks are required: a positive blackhole delivery difference versus active proofs with Holm p<0.05; a one-sided 95% delivery lower bound above −1 point versus frozen prior; a critical-violation upper bound below +1 point; and an energy-efficiency geometric-ratio upper bound below 1.02. Table 5 reports each check. These margins define an incremental replacement criterion. They do not replace the original absolute class-0 deadline-violation budget of 1.5%.')
h('IV. Results')
h('A. Delivery and matched controls',2)
p('Under blackholes, censored passive attention delivered 75.94% of offered packets, compared with 71.28% for active proofs. The paired difference was +4.67 percentage points, with 95% CI [2.36, 6.97] and Holm p=0.0161. Figure 4 shows setting-level intervals, while Fig. 5 retains every paired fixture and all blackhole comparator contrasts. No other setting’s delivery contrast against active proofs survived the 60-test correction. The effect therefore belongs to the specified blackhole condition rather than to all tested traffic regimes.')
p('Table 4 reports the full blackhole endpoint contrasts. Censored minus naive passive delivery was −0.17 points, and censored minus frozen prior was +0.06 points; neither contrast was significant after correction. The capacity matched MLP contrast was +0.44 points and also inconclusive. These controls rule out a claim that the observed replacement gain establishes a unique benefit from censoring, evidence updating or attention. They do not prove equivalence outside the declared margins.')
h('B. Energy efficiency and control traffic',2)
p('Blackhole energy per delivered payload byte decreased by 0.0006443 J/byte relative to active proofs, with 95% CI [−0.0009740, −0.0003145] and Holm p=0.0239. Total sensor radio energy instead rose from 22.546943 to 22.547473 J. This small descriptive increase, 0.000531 J, must not be reported as energy saving. The receive/listen state accounts for a mean 99.36% of modeled sensor radio energy across all 750 runs.')
p('Figure 6 separates the absolute energy numerator, delivered-byte efficiency and PHY transmission count. In blackhole settings, censored passive generated 60.57 fewer PHY transmissions per run on average than active proofs. Every passive/frozen arm requested zero forwarding proofs, while the active arm requested 3,369 across 150 runs. Beacons and existing reverse receipts remain in every arm. Because assumed receive/listen current exceeds transmit current, fewer transmitted frames need not reduce common-horizon joules in this always-listening model. Higher delivery explains the favorable energy-per-byte ratio.')
h('C. Gate outcome and service limits',2)
p('The three blackhole frozen-prior margins passed: the one-sided delivery lower bound was −0.405 percentage points, the critical-violation upper bound +0.925 points, and the geometric energy-efficiency ratio upper bound 1.00505. Together with the corrected active-proof delivery contrast, all four gate checks passed (Table 5). The upper critical bound is close to its +1-point margin, so the result should be read at its declared threshold rather than as a large safety reserve.')
p('The critical-violation decrease against active proofs did not survive multiple-testing correction. Absolute critical violations for censored passive remain approximately 26.2% under blackholes, far above the original 1.5% budget. Passing the replacement gate therefore establishes neither absolute deadline compliance nor a guaranteed quality-of-service controller.')
h('D. Implementation and trace audit',2)
p('The action/ledger verifier passed 67,522 decisions, 24,747 passive observations, 28,353 evidence updates and 75,750 node-energy identities. An additional endpoint audit reconstructed all 750 final delivery and critical-violation values from source decision times and sink reception events, checked beacon and decision counts, and reconciled PHY-state energy. All 30 gzip streams read successfully to end of stream. The full 51-test suite and the physical callback control passed again during this review.')
p('The trace of fixture 77004 was truncated during release-integrity checking and regenerated under the unchanged frozen implementation. All 750 canonical summary rows matched the pre-recovery summary exactly; no parameter, checkpoint or analysis rule changed. Audit code, provenance and the recovery record are distributed with the raw traces. This internal consistency check is not an independent radio implementation, external capture or a replay of every beacon-receive callback and service-state transition.')
h('V. Discussion and conclusions')
h('A. Critical vulnerabilities and limitations',2)
p('The strongest supported result is a replacement of the tested active proof mechanism. Its comparator changes feedback eligibility and expiry conventions as well as radio traffic, so causation cannot be assigned solely to proof airtime. The naive and frozen-prior controls provide no corrected evidence for a unique censoring effect. The earlier unsuccessful architecture comparisons remain part of the evidence record.')
p('State locality is incomplete. Geometry, packet history and reverse receipt paths are simulator inputs; beacon values use receive-gated metadata. Source-only service debt does not establish transit-aware distributed queue control. Forged receipts, deceptive telemetry, compromised sources, collusion and mobility are untested. Successful passive transmission is evidence of an attempt, not intended-receiver delivery, and the effective-count score is not a live intrusion detector.')
p('The four-second horizon, common central sink and 60-packet burst do not establish steady-state scalability or network lifetime. Candidate scoring is linear in the bounded candidate count after cache sorting, but each expiration poll scans every retained watch. With W records, memory is O(W) and expiration is O(W) per millisecond poll. There is no long-horizon reclamation policy or measured embedded inference cost. Always-listening radio assumptions also exclude the cost of waking specifically to obtain evidence.')
p('The passive strength 0.5, 150 ms expiry, 0.75 s cache life, current model and fixed inference checkpoints were not varied in the new holdout. Their sensitivity remains unknown. Deployment changes to state serialization, acknowledgment authentication, source-service tracking or ledger retirement must be assessed in a separate frozen experiment rather than inserted into this one after observing its outcome.')
h('B. High novelty opportunities',2)
p('A stronger next contribution would measure when evidence is worth acquiring. An adaptive observation policy could choose between silence, passive listening and authenticated probes using the expected reduction in delivery loss minus measured radio and processing cost. The discriminating experiment would match feedback semantics and data airtime, introduce duty cycling and adversarial receipts, and reserve new topology and hardware settings. This direction is not implemented in the present study; it follows from the attribution limits exposed here.')
h('C. Conclusions',2)
p('Receipt assisted passive feedback passed a prespecified incremental gate on 30 new simulation fixtures. Blackhole delivery improved by 4.67 percentage points over active proofs, and the declared frozen-prior non-inferiority margins held. The matched controls did not establish a unique censoring or attention advantage. Energy efficiency improved through delivered-byte yield while total modeled radio joules increased slightly. The defensible contribution is a reproducible feedback comparison with explicit attribution boundaries; authentication, embedded state exchange and absolute service compliance remain open.')
h('Acknowledgments')
p('No external funding was received. The authors declare no competing interests. This simulation study involved no human participants, clinical data or animal experiments.')
p('Data and code availability: code, frozen protocol, checkpoints, historical negative outcomes, raw traces and figure sources are available at https://github.com/DrShivamBhardwaj/JICCE_LA_CPPO_Reproducibility. The pre-outcome source freeze is 4027f071322cb5682f39329f656d76df3b68ff62. The complete earlier evidence manuscript is retained in publication/censored_feedback/.')
p('Author contributions: Shivam Bhardwaj, conceptualization, methodology, software, investigation and original draft; Jitendra Nath Srivastava, validation and review; Akash Sanghi, supervision and review. These roles retain the author-provided source statement. Generative AI assisted drafting, code preparation and figure generation; the authors are responsible for verification and final approval.')
h('References')
# All literature citations have appeared before the end exhibits.
assert len(order)==len(REFS),(len(order),set(REFS)-set(order))
for key,num in order.items():
 x=p('['+str(num)+'] '+REFS[key][0]);x.paragraph_format.line_spacing=1.15;x.paragraph_format.space_after=Pt(7)
# End exhibits follow journal review order; all are cited in body.
page();h('Algorithm 1 Observer feedback and routing')
for line in ['Input: received MAC events, cached telemetry, fixed policy θ, ledger mode and 150 ms timeout.','Output: next-hop requests and updated observer evidence coordinates.','Steps:','1. At a routing event, reject a visited or exhausted node and enforce the 22-hop cap.','2. Construct receive-gated candidates, retain at most eight and evaluate the masked fixed policy.','3. Register the observer–relay–packet record with expected link success; send the data frame.','4. On own first-hop ACK success, mark the record eligible; credit any stored uncredited observation.','5. On ACK failure, close the passive record without success or failure credit.','6. On a matched received forward attempt, retain maximum strength 0.5.','7. On a matched received reverse delivery receipt, retain maximum strength 1.0.','8. Apply a positive update only if eligible and strength exceeds credited strength; deduplicate repeats.','9. At each 1 ms poll, close expired records; censor missing evidence in the proposed arm.','10. Allow a later receipt to upgrade an acknowledged censored record; keep failure count at its prior.','11. Update source service debt on receipt or service expiry; retain the frozen policy for later decisions.']:
 x=p(line);x.paragraph_format.line_spacing=1.3;x.paragraph_format.space_after=Pt(5)
# References in tables reuse assigned numbers.
table(1,'Literature and experimental focus',['Study','Verified mechanism','Relation to this experiment'],[
['Watchdog {{watchdog}}','Overheard forwarding attempts','Established passive observation; reception ambiguity remains'],['2ACK {{ack}}','Reverse acknowledgment scheme','Explicit feedback adds traffic; no numerical cross-study ranking'],['TARM {{edge}}','Edge assisted trust routing','Trust-aware routing context; different state and deployment'],['PathGuard {{path}}','Trust and optimizer next-hop selection','Publisher abstract verified; not independently replicated'],['SMART {{smart}}','Learned multi-feature trust estimation','Detection task differs from forwarding evidence preference'],['FRDR {{frdr}} and RLFS-OR {{rlfs}}','Failure-aware and opportunistic learning','Different radio and duty-cycle assumptions'],['TG-RL-SR {{tgrl}} and secure DRL {{drlsecure}}','Trust gated and secure routing','Recent context; claims not treated as matched experimental effects'],['This study','ACK gated passive evidence; four controls','Frozen 30-fixture evaluation with retained negative stages']],[3.2,5.0,7.8])
table(2,'Frozen network settings',['Parameter','Value'],[
['Sensors and sink','100 sensors; one central sink'],['Field and range','150 × 150 m; 50 m hard range'],['Radio engine','ns-3.44; unslotted IEEE 802.15.4'],['Propagation','Log distance exponent 3; reference loss 40 or 54 dB'],['Run and offered load','4 s; 60 packets; 0.4 s warm-up'],['Source intervals','20 ms normal; 2 ms congested'],['Payload and identifiers','52 and 12 bytes; 64-byte application frame'],['Telemetry broadcast and cache','0.5 s period; 3 ms/node staggering; 0.75 s expiry'],['Feedback expiry and poll','150 ms; 1 ms'],['Traffic classes','Probabilities 0.35, 0.45, 0.20; deadlines 90, 350, 1,500 ms'],['Blackhole ratio','0 or 25%; non-forging forwarding faults'],['Assumed voltage and currents','3 V; TX 17.4 mA; RX/listen 18.8 mA; off 20 µA'],['Policy parameters','Attention 63,368; MLP 63,554; fixed weights'],['Holdout','30 fixtures × 5 settings × 5 arms = 750 runs']],[5.2,10.8])
table(3,'Feedback controls',['Arm','Evidence and expiry','Policy and interpretation'],[
['Active proofs','Separate proofs; legacy timeout and ACK eligibility','Attention; joint feedback and traffic comparator'],['Naive passive','Overhear/receipt positives; missing ACKed forwarding penalized','Attention; missing-observation control'],['Frozen prior','Same initial risk coordinates; no evidence updates','Attention; score-update control'],['Capacity MLP','Censored ledger with same radio observations','63554 parameters; representation control'],['Censored passive','Own ACK required; max-strength credit; missing evidence censored','63368-parameter attention; proposed replacement']],[3.2,6.4,6.4])
black=[]
for comp,label in [('active_proof','Active proofs'),('naive_passive','Naive passive'),('frozen_prior','Frozen prior'),('capacity_mlp_censored','Capacity MLP')]:
 vals=[]
 for endpoint,scale in [('delivery',100),('critical_violation',100),('energy_per_delivered_byte',1000)]:
  z=next(x for x in A if x['setting']=='blackhole' and x['comparator']==comp and x['endpoint']==endpoint);lo,hi=np.array(z['ci95'])*scale;vals.append(f"{z['difference']*scale:+.3f}\n[{lo:+.3f}, {hi:+.3f}]\np={z['holm_p']:.4g}")
 black.append([label,*vals])
table(4,'Blackhole paired effects',['Comparator','Delivery points','Critical violation points','Energy mJ per byte'],black,[3.0,4.1,4.4,4.5])
p('Values are censored passive minus comparator; intervals are unadjusted paired 95% intervals. p values are Holm adjusted over the original 60 tests. Negative critical-violation and energy-per-byte differences favor censored passive. These are differences, not absolute arm means.').paragraph_format.line_spacing=1.15
table(5,'Prespecified incremental gate',['Check','Frozen criterion','Observed result'],[
['Delivery versus active proofs','Positive gain and Holm p<0.05','+4.667 points; p=0.0161; pass'],['Delivery versus frozen prior','One-sided lower bound >−1 point','−0.405 points; pass'],['Critical violations versus frozen prior','One-sided upper bound <+1 point','+0.925 points; pass'],['Energy efficiency versus frozen prior','One-sided geometric upper ratio <1.02','1.00505; pass'],['Joint incremental decision','All four checks required','Pass; absolute QoS compliance not established']],[4.9,5.6,5.5])
for num,name,caption in [(1,'fig1_architecture','Architecture and state boundaries.'),(2,'fig2_algorithm_flow','Censored feedback flow.'),(3,'fig3_evaluation_flow','Frozen evaluation flow.'),(4,'fig4_delivery_effects','Delivery effects with paired 95% intervals.'),(5,'fig5_blackhole_estimation','Blackhole fixture outcomes and paired effects.'),(6,'fig6_energy_attribution','Radio effects with paired 95% intervals.')]:fig(num,name,caption)
# Short contextual notes define symbols without inflating captions.
# Full graph descriptions are already in paragraphs, legends and axis labels.
# Remove template ornament and theme-font overrides; scientific review style.
for st in d.styles:
 for node in list(st.element.iter(qn('w:pBdr'))):node.getparent().remove(node)
for para in d.paragraphs:
 for node in list(para._p.iter(qn('w:pBdr'))):node.getparent().remove(node)
 for run in para.runs:
  run.font.name='Times New Roman';run.font.color.rgb=RGBColor(0,0,0)
  rf=run._r.get_or_add_rPr().get_or_add_rFonts()
  for key in list(rf.attrib):
   if key.endswith('Theme'):del rf.attrib[key]
 if para.style.name=='Title':para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.runs[0].bold=True
# Comma-delimit adjacent numbered citations.
for para in d.paragraphs:
 for run in para.runs:
  if '][' in run.text:run.text=re.sub(r'\]\[',', ',run.text)
dest=OUT/'JICCE_Redesigned_Manuscript.docx';d.save(dest);(OUT/'JICCE_Redesigned_Manuscript.md').write_text('\n'.join(md))
with zipfile.ZipFile(dest) as z:xml=etree.fromstring(z.read('word/document.xml'));native=len(xml.xpath('//m:oMath',namespaces={'m':M}))
assert native==n_eq==10,(native,n_eq)
text='\n'.join(x.text for x in d.paragraphs);assert all('Fig. '+str(i) in text for i in range(1,7));assert all('Table '+str(i) in text for i in range(1,6));assert 'PRISMA-inspired' in text
checks={'abstract_words':len(abstract.split()),'references':len(order),'native_equations':native,'figures':6,'tables':5,'algorithms':1,'all_exhibits_cited':True,'figure_formats':['svg','pdf','900dpi png'],'frozen_outcomes_changed':False,'submission_certified':False,'format':'JICCE review text: 9.5 pt, double-spaced A4, 2.5 cm margins; Arabic table numbers','author_side_items':['Corresponding author phone/fax details if required by portal','Final coauthor approval and forms','Editorial assessment of contribution and manuscript length']}
(OUT/'document_checks.json').write_text(json.dumps(checks,indent=2));(OUT/'reference_verification.json').write_text(json.dumps([dict(number=order[k],reference=v[0],url=v[1],verification=v[2]) for k,v in REFS.items()],indent=2));print(dest);print(checks)
