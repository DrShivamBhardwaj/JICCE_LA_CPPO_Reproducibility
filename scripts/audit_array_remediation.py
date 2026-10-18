"""Check new exhibits against completed data and preserve the retained manuscript."""
from pathlib import Path
import argparse,json,hashlib,zipfile
from lxml import etree
from docx import Document
from array_remediation_integration import primary_rows,TITLE
R=Path(__file__).resolve().parents[1];P=R/'publication/array'
def sha(b):return hashlib.sha256(b).hexdigest()
def main():
    ap=argparse.ArgumentParser();ap.add_argument('--baseline',type=Path,required=True);args=ap.parse_args()
    old=Document(args.baseline);new=Document(P/'Array_Manuscript.docx')
    assert len(old.tables)==11 and len(new.tables)==12
    table=lambda t:[[c.text for c in row.cells] for row in t.rows]
    assert [table(t) for t in old.tables]==[table(t) for t in new.tables[:11]]
    math=lambda d:[etree.tostring(e,method='c14n') for e in d._element.xpath('//m:oMath')]
    assert len(math(old))==10 and math(old)==math(new)
    with zipfile.ZipFile(args.baseline) as z:before={sha(z.read(n)) for n in z.namelist() if n.startswith('word/media/')}
    with zipfile.ZipFile(P/'Array_Manuscript.docx') as z:after={sha(z.read(n)) for n in z.namelist() if n.startswith('word/media/')}
    added={sha(p.read_bytes()) for p in (P/'remediation_figures').glob('fig*.png')}
    assert len(before)==12 and len(after)==14 and before<=after and after-before==added
    edits=json.loads((P/'remediation_edits.json').read_text());mapping={e['before']:e['after'] for e in edits['replacements']};texts=[p.text for p in new.paragraphs]
    for p in old.paragraphs:
        if p.text:assert mapping.get(p.text,p.text) in texts,'Unmapped change: '+p.text[:120]
    for t in edits['added_paragraphs']:
        if t:assert t in texts
    analysis=json.loads((R/'outputs/remediation_v1/analysis.json').read_text())
    assert table(new.tables[-1])[1:]==primary_rows(analysis)
    assert TITLE in texts and len(new.inline_shapes)==14
    result={'status':'passed','original_tables_preserved':11,'native_equations_preserved':10,'original_images_preserved':12,'new_images_verified':2,'primary_contrasts_reconciled':6,'baseline_sha256':sha(args.baseline.read_bytes()),'final_docx_sha256':sha((P/'Array_Manuscript.docx').read_bytes()),'physical_validation':False,'original_signed_deadline_complete':False}
    (P/'remediation_manuscript_audit.json').write_text(json.dumps(result,indent=2)+'\n');print(json.dumps(result,indent=2))
if __name__=='__main__':main()
