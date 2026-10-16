"""Append locked holdout evidence while preserving all existing native equations."""
from pathlib import Path
from copy import deepcopy
import json,zipfile
from lxml import etree
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.text import WD_BREAK,WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1];out=ROOT/'publication/censored_feedback';out.mkdir(parents=True,exist_ok=True)
source=ROOT/'publication/ns3_evidence/JICCE_LA_CPPO_Telemetry_Review_Draft.docx';d=Document(source);a=json.loads((ROOT/'outputs/ns3_censored/paired_analysis.json').read_text());decision=json.loads((ROOT/'outputs/ns3_censored/decision.json').read_text());v=json.loads((ROOT/'outputs/ns3_censored/verification.json').read_text());assert v['status']=='passed' and v['runs']==750
anchor=next(p for p in d.paragraphs if p.text=='X. Discussion')
def add(text,style='Normal'):
 p=anchor.insert_paragraph_before(text);p.style=style
 for r in p.runs:r.font.name='Times New Roman'
 return p
M='http://schemas.openxmlformats.org/officeDocument/2006/math'
def mr(text):
 r=etree.Element('{'+M+'}r');t=etree.SubElement(r,'{'+M+'}t');t.text=text;return r
def sub(base,index):
 x=etree.Element('{'+M+'}sSub');e=etree.SubElement(x,'{'+M+'}e');e.append(mr(base));s=etree.SubElement(x,'{'+M+'}sub');s.append(mr(index));return x
def equation(items,number):
 p=add('');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;math=etree.SubElement(p._p,'{'+M+'}oMath')
 for item in items:math.append(mr(item) if isinstance(item,str) else item)
 math.append(mr('    ('+str(number)+')'))
add('G. Receipt assisted censoring and locked holdout validation','Heading 2')
add('The 76001–76030 telemetry fixtures informed one feedback redesign and were not reused for its inference. The new ledger combines passive forwarding observations with delivery receipts already present in the network. Overhearing is established watchdog practice and has known collision, receiver-loss and collusion ambiguities [28]; no invention of watchdog monitoring is claimed. The tested change conditions credit on the observer’s own successful MAC acknowledgment, censors missing evidence, and deduplicates repeated observations without generating separate proof packets. A received forwarding frame supplies half-strength evidence of an attempt; a reverse delivery receipt supplies full-strength evidence of path delivery. Neither observation is an authenticated attack label.')
add('For observer u, selected relay v and packet p, the ledger retains the strongest observed evidence s and the already credited amount c. A passive frame and a delivery receipt cannot accumulate beyond one unit for the same record. The evidence-strength rule is')
equation([sub('s','uvp'),' ← max{',sub('s','uvp'),', 0.5 ',sub('I','overheard'),', ',sub('I','receipt'),'}'],11)
add('Only after the first-hop MAC acknowledgment succeeds and the uncredited evidence is positive, the observer increases the relay’s effective trust count using its expected link-success weight q:')
equation([sub('α','uv'),' ← ',sub('α','uv'),' + clip(',sub('q','uv'),' max{0, ',sub('s','uvp'),' − ',sub('c','uvp'),'}, 0.05, 1),    ',sub('c','uvp'),' ← ',sub('s','uvp')],12)
add('The corresponding failure count remains at its prior in the censored arm. Failed first links and missing observations do not increment it. A late receipt can still resolve an acknowledged but censored record. These effective counts supply the existing routing score; their asymmetric observation process does not define a calibrated probability of maliciousness. The rate-anomaly term is inactive in this transfer. The naive passive control uses the same evidence and deduplication but penalizes missing acknowledged forwarding, whereas the active-proof control retains the earlier proof and timeout convention.')
add('Design, implementation and decision criteria were published at commit 4027f071322cb5682f39329f656d76df3b68ff62 before evaluating the new fixtures. Seeds 77001–77030 pair with the unchanged checkpoints trained on 41001–41030; technical pilot 76901 is excluded. Five arms cross the same five radio/traffic settings, producing 750 runs and 45,000 source packets. The additional controls are frozen-prior attention, which retains the same initial risk coordinates without updates, and the capacity-matched MLP with the censored ledger. All arms retain the existing beacon, reverse-receipt and listening costs. No extra listening mode is enabled; an actual duty-cycled implementation may incur different acquisition and CPU costs.')
add('The 60 two-sided paired tests form one Holm family. The all-required mechanism gate combines a positive, Holm-significant blackhole delivery contrast against active proofs with three one-sided 95% non-inferiority checks against frozen-prior attention: delivery above −1 percentage point, critical violations below +1 point, and geometric energy per delivered byte below a 1.02 ratio. This is a bounded incremental criterion, not the original absolute QoS constraint or a field-deployment criterion.')
add('Under blackholes, censored passive attention delivered 75.94% of packets and active-proof attention delivered 71.28%. The paired gain was 4.67 percentage points (95% CI 2.36–6.97; Holm p=0.0161). Radio energy per delivered payload byte decreased by 0.0006443 J/B (95% CI −0.0009740 to −0.0003145; Holm p=0.0239). The critical-violation decrease of 4.52 points did not survive the 60-test correction. No other setting’s delivery contrast against active proofs survived that correction. Figure 8 shows the unadjusted delivery intervals; Table XIII reports the blackhole contrasts.')
add('The frozen-prior comparison changed blackhole delivery by only +0.056 points (95% CI −0.498 to +0.609; Holm p=1). Its one-sided lower delivery bound was −0.405 points, the upper critical-violation bound was +0.925 points, and the upper geometric energy ratio was 1.00505. All frozen gate checks passed. Neither the naive passive control nor the censored capacity MLP was significantly worse than censored attention on the blackhole endpoints. The result supports replacing the costly proof-based mechanism under the specified conditions; it does not demonstrate a unique performance benefit from censoring, from updating evidence, or from attention itself.')
add('All four passive or frozen-prior arms emitted zero forwarding-proof packets; the active arm requested 3,369 across its 150 runs. The audit reconstructed 67,522 feature/action decisions, 24,747 matched passive observations, 28,353 evidence updates and 75,750 node-energy identities. It checked the selected relay, observer range, data-frame provenance, local risk counts and per-record deduplication. These internal traces do not provide external packet-capture, authenticated-feedback or IDS validation. Absolute critical-class violations remain about 26% in the blackhole setting, far above the original 1.5% budget.')
abstract=next(p for p in d.paragraphs if p.text.startswith('Routing under deadlines depends'))
abstract.text='Routing feedback can change both forwarding decisions and the traffic they depend on. We evaluate local-attention constrained PPO routing through capacity-matched comparisons and packet-level wireless transfer. A 600-run telemetry study found no corrected architecture advantage, while the proof-based mechanism reduced delivery relative to a proof-disabled control. We then froze a receipt-assisted passive evidence ledger and evaluated 750 runs on 30 new fixtures. Under blackholes, censored feedback delivered 75.94% of packets versus 71.28% with active proofs, a gain of 4.67 percentage points (95% CI 2.36–6.97; Holm p=0.0161). The prespecified incremental gate passed, including non-inferiority to frozen-prior attention. Censoring and attention did not independently establish superiority over their matched controls. The evidence supports a bounded replacement of costly proof traffic. It does not establish live intrusion detection, hardware energy savings, general routing superiority or satisfaction of the original absolute service budgets.'
conclusion=next(p for p in d.paragraphs if p.text.startswith('LA-CPPO combines candidate scoring'))
conclusion.text='Historical and capacity-matched evaluations did not establish the proposed architecture’s joint superiority. Explicit radio telemetry then exposed a practical cost of active forwarding proofs. A separately frozen receipt-assisted passive ledger passed its incremental holdout gate: blackhole delivery improved over active proofs and remained within the declared delivery, critical-violation and energy margins against frozen-prior attention. The matched naive and frozen-prior controls prevent attribution of this gain specifically to censoring or evidence updates, and the architecture comparison remains inconclusive. The defensible result is a bounded feedback-mechanism replacement under the evaluated radio and traffic assumptions. Authentication, live IDS, physical energy calibration and the original absolute QoS guarantees remain unestablished.'
for p in d.paragraphs:
 if 'both ns-3 transfer programs' in p.text:p.text=p.text.replace('both ns-3 transfer programs','all ns-3 validation programs and the passive-feedback ledger')
add('Beacon energy and queue values are carried in simulator-side metadata keyed to a transmitted beacon identifier and become available only after an actual radio receive event. Their values are not decoded from serialized beacon fields. The charged 64-byte beacon frame is therefore a modeled telemetry transport, not a demonstrated embedded wire implementation. Geometry and packet-class configuration are also simulator inputs; this transfer does not establish a distributed embedded deployment.')
# Insert the new reference after the existing bibliography, before end tables.
ref=next(p for p in d.paragraphs if p.text.startswith('[27]'))
x=OxmlElement('w:p');ref._p.addnext(x)
from docx.text.paragraph import Paragraph
p=Paragraph(x,ref._parent);p.style=ref.style;p.add_run('[28] S. Marti, T. J. Giuli, K. Lai, and M. Baker, “Mitigating routing misbehavior in mobile ad hoc networks,” in Proceedings of MobiCom, 2000. DOI: 10.1145/345910.345955.').font.name='Times New Roman'
# One compact end table; journal layout keeps figures and tables at the end.
d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
d.add_paragraph('Table XIII. Blackhole holdout effects for censored passive attention minus each comparator.')
t=d.add_table(rows=1,cols=4)
for c,text in zip(t.rows[0].cells,['Comparator','Delivery pp','Critical violations pp','Radio energy J/B']):c.text=text
for comp,label in [('naive_passive','Naive passive'),('active_proof','Active proofs'),('frozen_prior','Frozen prior'),('capacity_mlp_censored','Capacity MLP')]:
 cells=t.add_row().cells;cells[0].text=label
 for cell,endpoint,scale in zip(cells[1:],['delivery','critical_violation','energy_per_delivered_byte'],[100,100,1]):
  z=next(x for x in a if x['setting']=='blackhole' and x['comparator']==comp and x['endpoint']==endpoint)
  cell.text=f"{scale*z['difference']:+.4g}\n[{scale*z['ci95'][0]:+.4g}, {scale*z['ci95'][1]:+.4g}]\np={z['holm_p']:.4g}"
t.autofit=False
for i,row in enumerate(t.rows):
 for cell in row.cells:
  cell.width=Cm(4)
  for p in cell.paragraphs:
   p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(4)
   for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(9);r.bold=i==0
borders=d.tables[0]._tbl.tblPr.find(qn('w:tblBorders'))
if borders is not None:t._tbl.tblPr.append(deepcopy(borders))
d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(out/'figures/fig8.png'),width=Cm(15))
d.add_paragraph('Fig. 8. Holdout delivery effects with unadjusted paired 95% intervals. Positive differences favor censored passive attention.')
dest=out/'JICCE_Censored_Feedback_Review_Draft.docx';d.save(dest)
def math_items(path):
 with zipfile.ZipFile(path) as z:r=etree.fromstring(z.read('word/document.xml'));return [etree.tostring(x) for x in r.xpath('//m:oMath',namespaces={'m':M})]
old=math_items(source);new=math_items(dest);assert all(x in new for x in old) and len(new)==len(old)+2
(out/'document_checks.json').write_text(json.dumps({'old_native_equations_preserved':len(old),'native_equations_total':len(new),'holdout_runs':v['runs'],'incremental_gate_pass':decision['pass_all'],'submission_certified':False},indent=2));print(dest)
