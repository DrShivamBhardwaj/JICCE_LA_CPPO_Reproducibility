"""Array adaptation retaining native equations and the frozen evidence record."""
from pathlib import Path
from copy import deepcopy
import json,re,zipfile
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1];out=ROOT/'publication/array';out.mkdir(exist_ok=True)
d=Document(ROOT/'publication/redesigned/JICCE_Redesigned_Manuscript.docx')
analysis=json.loads((ROOT/'outputs/auth_policy/analysis.json').read_text())
audit=json.loads((ROOT/'outputs/auth_policy/audit.json').read_text())
virtual=json.loads((ROOT/'outputs/virtual_boards/verification.json').read_text())
title='Receipt Assisted Feedback and Authenticated Policy Transfer for Wireless Sensor Routing'
abstract=('Routing feedback consumes radio capacity and may confuse missing evidence with forwarding failure. '
'We evaluate receipt assisted feedback and a separate authenticated transport for frozen sensor-routing policies. '
'The first experiment fixes the feedback ledger before 750 IEEE 802.15.4 simulation runs on 30 new fixtures, '
'with active proof, naive passive, frozen prior and capacity matched policy controls. Under blackholes, delivery '
'increases from 71.28% to 75.94% relative to active proofs, a paired gain of 4.67 percentage points '
'(95% interval 2.36–6.97; Holm adjusted p=0.0161). Energy per delivered byte improves while total modeled '
'radio energy increases slightly. A second protocol evaluates attention, multilayer perceptron and greedy '
'routing in 180 runs using byte-serialized neighbor state, hop authentication, replay protection and signed '
'sink receipts. Attention delivery is 81.61% in benign conditions and 61.11% with selective dropping; '
'paired comparisons establish no superiority over the controls. Four virtual serial devices also exercise '
'the host radio adapter and reveal a receive-time freshness bug, which is corrected. The corrected integration '
'delivers all ten offered packets in that functional fixture. A separate four-board Cortex-M3 emulator executes compiled radio-interface firmware; all three host-policy controls deliver ten of ten packets. These results provide traceable feedback and '
'implementation evidence while exposing transfer and attribution limits. Virtual serial interfaces and '
'modeled radio energy do not establish physical board operation, measured energy consumption or absolute deadline compliance.')
for p in d.paragraphs:
 if p.style.name=='Title':p.text=title
 if p.text.startswith('Running title:'):p.text=''
 if p.text.startswith('Routing feedback consumes the channel'):p.text=abstract
 if p.text.startswith('Index Terms:'):p.text='Keywords: Wireless sensor networks; Routing feedback; Authentication; Policy transfer; Reproducibility; Virtual validation'
 mapping={'I. Introduction':'1. Introduction','A. Contribution and research questions':'1.1 Contribution and research questions',
 'II. System model and assumptions':'2. System model and assumptions','III. Receipt assisted passive feedback':'3. Receipt assisted passive feedback',
 'IV. Experimental results':'4. Experimental results','V. Discussion and conclusions':'6. Discussion',
 'A. Critical vulnerabilities and limitations':'6.1 Critical vulnerabilities and limitations',
 'B. High novelty opportunities':'6.2 High novelty opportunities','D. Conclusions':'7. Conclusions'}
 if p.text in mapping:
  old=p.text;p.text=mapping[old]
  if old=='D. Conclusions':p.style='Heading 1'
# Convert any remaining Roman main sections and lettered subsections.
section=0;sub=0
for p in d.paragraphs:
 if p.style.name=='Heading 1':
  for roman,num in [('II.',2),('III.',3),('IV.',4)]:
   if p.text.startswith(roman+' '):p.text=str(num)+'. '+p.text[len(roman)+1:]
  m=re.match(r'([1-7])\. ',p.text)
  if m:section=int(m[1]);sub=0
 elif p.style.name=='Heading 2' and re.match(r'[A-D]\. ',p.text):
  sub+=1;p.text=f'{section}.{sub} '+p.text[3:]
# Replace the earlier development-only extension with integrated evaluated evidence.
remove=False
for p in list(d.paragraphs):
 if 'Post holdout engineering extension' in p.text:remove=True
 if p.text=='7. Conclusions':remove=False
 if remove:p._element.getparent().remove(p._element)
anchor=next(p for p in d.paragraphs if p.text=='6. Discussion')
model_anchor=next(p for p in d.paragraphs if p.text=='2. System model and methods')
model_anchor.insert_paragraph_before('A separate implementation stage asks whether the same frozen policies can operate on authenticated wire state and node-local transit service. It evaluates 180 new radio runs and a virtual serial-interface fixture without transferring the earlier performance claim to the new transport. Sections 2–4 describe the frozen feedback benchmark; Section 5 defines and evaluates the authenticated branch.')
def before(text,style=None):return anchor.insert_paragraph_before(text,style)
before('5. Authenticated transport and virtual validation','Heading 1')
before('5.1 Local state and authenticated evidence','Heading 2')
before('The transport in Fig. 7 is a separate implementation branch. It serializes coordinates, modeled remaining radio energy, queue state and transmit time into 58-byte application telemetry frames. Each node uses only its configured own and sink coordinates, its received authenticated cache, local MAC outcomes and packet-carried history. Pairwise HMAC-SHA256 with a 128-bit tag protects each hop [20]. Provisioned 64-bit boot sessions and a 64-entry replay window reject stale sessions and duplicates. Data binds immutable source fields with a source authenticator. The sink signs the source, session and packet identity using Ed25519 [21]; a signed receipt returns along node-local reverse records. Maximum application frame size is 114 bytes.')
before('Passive third-party overhearing receives no authenticated forwarding credit in this branch because pairwise keys cannot verify another pair’s onward message. Local successful MAC acknowledgment and a verified sink receipt jointly enable positive forwarding credit. Missing receipts are censored. Node records are bounded at 128 entries with 2 s expiry. Keys and sessions are provisioned out of band; protected key storage, online key exchange and truthful telemetry from a compromised key holder are not provided.')
before('5.2 Frozen policy transfer protocol','Heading 2')
before('The protocol and source were committed before observing the new outcomes (commit 5f06b9213e73206aad11ec8f3a96079a17011a20). Thirty independent topology seeds, 83001–83030, each pair one frozen attention checkpoint with its capacity matched multilayer perceptron checkpoint and a greedy progress control. A six-node diamond uses two first-hop alternatives and two downstream relays. Nominal coordinates are (0,0), (25,−15), (25,15), (50,−15), (50,15) and (75,0) m, with independent ±2 m jitter. Range is 38 m. Benign and selective-dropping conditions produce 180 runs. In the latter, relay 1 discards received data with an even packet identifier after MAC reception; this label is unavailable to the router.')
before('Each 7.5 s run offers 60 packets with 16-byte payloads, beginning at 0.5 s at 75 ms intervals, with traffic class cycling through the three deadline classes. Authenticated unicast telemetry is attempted to every provisioned peer every 0.5 s with 2 ms staggering. Out-of-range requests and all modeled radio costs remain in the record. The three arms share strict distance-progress filtering and the eight-candidate limit. Greedy selects the shortest remaining sink distance; learned arms retain their original weights. Fixed ns-3 seeds and streams pair fixtures, although different actions consume different random sequences.')
before('The 93-element observation comprises 13 context coordinates and eight ten-coordinate candidates. Neighbor geometry and energy come from decoded telemetry. MAC success uses a Beta(9,1) prior; acknowledgment delay uses an exponentially weighted estimate with coefficient 0.2 and a 5 ms initial value. Received LQI replaces the training RSSI proxy. Queue coordinates use raw MAC pending-request counts rather than the normalized pressure used during training. Forwarding risk is a positive-only transform of signed-receipt success counts and current link reliability. Each node tracks separate class delay and loss debts. These declared substitutions constitute distribution shift; the transfer does not reproduce the training observation model or establish calibrated intrusion probabilities.')
before('5.3 Transfer results and audit','Heading 2')
before('Table 6 and Fig. 8 report all 180 runs. Attention delivers 81.61% in benign conditions, compared with 82.72% for greedy and 83.72% for the multilayer perceptron. Its paired delivery differences are −1.11 points versus greedy (95% interval −3.39 to 0.83) and −2.11 points versus the perceptron (−4.50 to 0.28). Under selective dropping, attention delivers 61.11%, greedy 62.89% and the perceptron 60.50%; the attention differences are −1.78 points (−6.50 to 2.67) and +0.61 points (−2.56 to 3.78), respectively. Intervals use 10,000 paired percentile bootstrap samples across fixtures with seed 94001. These exploratory intervals are not multiplicity-adjusted superiority tests.')
before('The transfer establishes execution with byte-serialized simulated radio frames and multiple candidate choices, but supplies no evidence of an attention advantage. Critical deadline failure remains 31.17% in benign attention runs and 47.83% under selective dropping, counting undelivered packets as failures. Total modeled radio energy is approximately 2.519 and 2.520 J per six-node run. Always-on listening and repeated authenticated control traffic limit energy differentiation; no embedded CPU or cryptographic energy is included.')
before(f"The audit reconstructed {audit['decisions']:,} decisions, including {audit['multichoice']:,} with multiple candidates, checked {audit['received_frames']:,} received byte frames against prior transmissions, verified {audit['signatures']:,} public sink-signature occurrences and reconciled {audit['energy_identities']:,} radio-energy identities. It also matched candidate cache values to received beacon bytes and recomputed fixed-policy probabilities. Secret hop tags are checked by runtime nodes rather than exported for public re-verification. The original 750-run traces, checkpoints and gate remain unchanged.")
before('5.4 Virtual serial devices and firmware boundary','Heading 2')
before('Four pseudo-terminal serial devices implement the board-facing TX, RX, MAC-confirm and telemetry contract while ns-3 transports the exact application bytes. Four independent host-adapter processes exercise the same serial interface intended for later board integration. The first diagnostic delivered five of ten packets and exposed false freshness rejections: the host sampled its clock before a blocking serial read. Sampling again after receipt corrected the defect. The corrected functional fixture delivers ten of ten packets, settles service at both intermediate relays and records no freshness rejection. Both captures are retained; this single fixture is not a reliability estimate.')
before('The pseudo-terminal devices model the firmware API without MCU execution. Rechecks with greedy, multilayer perceptron and attention host policies each deliver ten of ten packets. A separate Renode 1.17.0 backend executes compiled C firmware on four emulated CC2538 ARM Cortex-M3 boards (Fig. 7). Firmware parses UART JSON, enforces frame bounds, drives the IEEE 802.15.4 radio FIFO and correlates MAC acknowledgments. Host processes retain routing inference and cryptography. All three controls deliver ten of ten packets and return ten verified source receipts each. Each run passes five probes: labeled uncalibrated telemetry, exact 114-byte delivery and acknowledgment, rejection of 115-byte frames, malformed-hex rejection and failed out-of-range transmission. The independent capture audit verifies received bytes against MCU UART output, public sink signatures and packet identities. Development captures are retained for startup synchronization and metadata corrections. A stricter later audit rejects an older greedy capture because its relay-1 final log is missing, despite its retained summary. New runs require all ten offered packet identities at the sink, ten runtime-accepted signed source receipts, complete source/relay final records, five successful probes and a public capture audit before receiving a pass status; all three new controls meet this gate.')
before('Renode uses a deterministic range medium, not the ns-3 physical-layer model of the statistical study. Its energy fraction is a fixed placeholder explicitly labeled emulated_no_power_model; queue telemetry counts pending MCU acknowledgment records. No battery result follows from this fixture. The four-node chain provides one feasible next hop per relay, so successful execution does not compare policy quality or estimate reliability. Deadline debts remain nonzero. Compiled firmware execution is demonstrated only in emulation; physical clock and pin initialization, embedded policy/crypto execution, physical RF behavior and calibrated current measurements remain unvalidated.')
# Update the conclusion and keep the original attribution limits.
for p in d.paragraphs:
 if 'Physical callback' in p.text:p.text=p.text.replace('Physical callback','Simulated radio callback')
 if 'physical callback control' in p.text:
  p.text=p.text.replace('The full 51-test suite and the physical callback control passed again during this review.','The frozen-stage 51-test suite and radio callback control passed; the expanded implementation suite now contains 83 passing tests.')
 if 'Forged receipts, deceptive telemetry' in p.text:
  p.text=p.text.replace('Forged receipts, deceptive telemetry','In that frozen benchmark, forged receipts, deceptive telemetry')
 if p.text.startswith('Receipt assisted passive feedback passed a prespecified'):
  p.text=('Receipt assisted passive feedback passed a prespecified incremental gate on 30 new simulation fixtures, improving blackhole delivery by 4.67 percentage points over active proofs. Matched controls did not establish a unique censoring or attention advantage. A separate authenticated transport executes frozen attention and capacity matched policies using received wire state and local service records. Its 180-run transfer study also establishes no attention superiority. Virtual serial devices expose and help correct a host receive-time defect. Together, these stages support a reproducible feedback comparison and an auditable implementation boundary. Compiled radio-interface firmware also executes on four emulated Cortex-M3 boards. Absolute service compliance, physical-board firmware operation, embedded policy/crypto execution and physical energy validation remain open.')
# Format as a readable single-column Elsevier review manuscript.
refs_anchor=next(p for p in d.paragraphs if p.text=='References')
refs_anchor.insert_paragraph_before('Declaration of generative AI use','Heading 1')
refs_anchor.insert_paragraph_before('ChatGPT assisted manuscript drafting, code preparation, analysis scripts and programmatic figure plotting. All numerical results reported here come from the retained executable experiments and traces. Responsibility for scientific verification, interpretation and final approval rests with the authors.')
# Apply the corresponding author's confirmed authorship and declarations.
meta=json.loads((out/'author_metadata.json').read_text())
for p in list(d.paragraphs):
 if p.text=='Shivam Bhardwaj, Jitendra Nath Srivastava and Akash Sanghi':
  p.text='Shivam Bhardwaj¹*, Dr. Abhishek Kumar Pandey² and Dr. Abhinandan Tripathi³'
 elif p.text.startswith('Faculty of Engineering and Technology, Invertis University'):
  p.text='¹ Assistant Professor, Department of Computer Applications (MCA), United Institute of Management, Prayagraj, Uttar Pradesh, India\n² '+meta['authors'][1]['affiliation']+'\n³ '+meta['authors'][2]['affiliation']
 elif p.text.startswith('Corresponding author:'):
  p.text='*Corresponding author: Shivam Bhardwaj; '+meta['authors'][0]['email']+'\nAbhishek Kumar Pandey: '+meta['authors'][1]['email']+'; ORCID: '+meta['authors'][1]['orcid']+'\nAbhinandan Tripathi: '+meta['authors'][2]['email']
 elif p.text.startswith('No external funding was received.'):
  p.text=meta['funding']+' '+meta['competing_interests']+' This simulation and emulation study involved no human participants, clinical data or animal experiments.'
 elif p.text.startswith('Author contributions:'):
  # Old authors' roles cannot be transferred to newly listed coauthors.
  p._element.getparent().remove(p._element)
for section in d.sections:
 section.top_margin=section.bottom_margin=Cm(2.5);section.left_margin=section.right_margin=Cm(2.5)
normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
for p in d.paragraphs:
 p.paragraph_format.line_spacing=1.15
 if p.style.name.startswith('Heading'):p.paragraph_format.keep_with_next=True
 for r in p.runs:
  r.font.name='Times New Roman';r.font.color.rgb=RGBColor(0,0,0)
  if p.style.name not in ['Title','Heading 1','Heading 2']:r.font.size=Pt(11)
# Additional exhibits follow the original numbered exhibits.
d.add_page_break();p=d.add_paragraph('Table 6. Authenticated policy transfer.');p.paragraph_format.keep_with_next=True
t=d.add_table(rows=1,cols=5)
heads=['Condition / policy','Delivery (%)','Receipts (%)','Critical failure (%)','Radio energy (J)']
for c,x in zip(t.rows[0].cells,heads):c.text=x
for setting in ['benign','selective_drop']:
 for arm in ['greedy','mlp','attention']:
  z=next(x for x in analysis['arm_means'] if x['setting']==setting and x['arm']==arm)
  vals=[('Benign' if setting=='benign' else 'Dropping')+' / '+arm,f"{z['delivery']*100:.2f}",f"{z['verified_receipts']*100:.2f}",f"{z['critical_violation']*100:.2f}",f"{z['total_radio_j']:.4f}"]
  for c,x in zip(t.add_row().cells,vals):c.text=x
borders=OxmlElement('w:tblBorders')
for side in ['top','bottom','left','right','insideH','insideV']:
 e=OxmlElement('w:'+side);e.set(qn('w:val'),'single' if side in ['top','bottom'] else 'nil');e.set(qn('w:sz'),'6');borders.append(e)
t._tbl.tblPr.append(borders)
for i,row in enumerate(t.rows):
 for j,c in enumerate(row.cells):
  c.width=Cm([5,2.7,2.7,3,2.6][j])
  for p in c.paragraphs:
   p.paragraph_format.space_before=p.paragraph_format.space_after=Pt(5)
   for r in p.runs:r.font.size=Pt(9);r.bold=i==0
d.add_paragraph('Means over 30 fixtures per condition and policy. Receipt rates count source-verified sink receipts. Critical failure includes undelivered critical packets. Energy covers modeled radio residence, not board or host CPU energy.')
for number,name,caption in [(7,'fig7_authenticated_virtual','Authenticated state and virtual interfaces.'),(8,'fig8_policy_transfer','Frozen policy transfer and paired effects.')]:
 d.add_page_break();p=d.add_paragraph();p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(out/'figures'/f'{name}.png'),width=Cm(16))
 d.add_paragraph(f'Fig. {number}. {caption}')
dest=out/'Array_Manuscript.docx';d.save(dest)
with zipfile.ZipFile(dest) as z:native=z.read('word/document.xml').count(b'<m:oMath>')
assert native==10
(out/'Array_Manuscript.md').write_text('\n\n'.join(p.text for p in d.paragraphs))
highlights=['A frozen 750-run study isolates feedback replacement effects.',
'Authenticated wire state supports a separate 180-run policy transfer study.',
'Attention shows no established delivery advantage over matched controls.',
'Virtual serial devices expose and correct a receive-time freshness defect.']
assert all(len(x)<=85 for x in highlights)
(out/'Highlights.txt').write_text('\n'.join(highlights)+'\n')
(out/'document_checks.json').write_text(json.dumps(dict(abstract_words=len(abstract.split()),native_equations=native,figures=8,tables=6,references=21,virtual_validation=True,physical_validation=False,submission_certified=False),indent=2))
print(dest)
