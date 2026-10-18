"""Integrate the independently audited paired Cooja evaluation without replacing old families."""
from pathlib import Path
from copy import deepcopy
import json
from docx.shared import Pt,Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
LABELS={'active':'Active proofs','naive':'Naive passive','frozen':'Frozen prior','censored_proof':'Censored + proofs','censored':'Censored passive'}
def table_rows(analysis):
 rows=[]
 for attack in [0,1]:
  for arm in ['active','naive','frozen','censored_proof','censored']:
   a=next(x for x in analysis['groups'] if (x['attack'],x['timing'],x['arm'])==(attack,0,arm));b=next(x for x in analysis['groups'] if (x['attack'],x['timing'],x['arm'])==(attack,1,arm))
   rows.append([('Benign' if not attack else 'Blackhole'),LABELS[arm],f"{100*a['delivery']:.2f}",f"{100*b['delivery']:.2f}",f"{100*b['critical_sink_deadline_failure']:.2f}",f"{b['total_tx_state_s']:.3f}"])
 return rows

def effect(x):
 return f"{100*x['mean_difference']:+.2f} percentage points (95% paired interval {100*x['ci95'][0]:+.2f} to {100*x['ci95'][1]:+.2f}; Holm-adjusted p={x['p_holm']:.4g})"

def apply(doc,root):
 root=Path(root);out=root/'publication/array';a=json.loads((root/'outputs/cooja_feedback/primary_600/analysis.json').read_text());assert a['status']=='passed_paired_seed_analysis' and a['runs']==600
 changes=[]
 for arm,value in [('active',79.11),('censored',32.94),('naive',54.89)]:
  group=next(x for x in a['groups'] if (x['attack'],x['timing'],x['arm'])==(1,1,arm));assert round(100*group['delivery'],2)==value
 abstract=doc.paragraphs[next(i for i,p in enumerate(doc.paragraphs) if p.text=='Abstract')+1];old=abstract.text
 abstract.text=old.replace('Serial-interface testing identifies and corrects a receive-time freshness error. Earlier UART-to-radio firmware runs on four emulated Cortex-M3 boards; neural inference and cryptography execute on the host PC. Each host-policy control delivers ten of ten packets.','In 600 Cooja transfer runs, a relay that starts dropping after warm-up exposes a failure: delivery with scheduled processing delays is 32.94% with censored feedback and 79.11% with active proofs.')
 changes.append(dict(before=old,after=abstract.text));assert len(abstract.text.split())<=250
 conclusion=next(p for p in doc.paragraphs if p.text.startswith('ACK-gated censored passive feedback passed the prespecified'))
 old=conclusion.text
 parts=['ACK-gated censored passive feedback improves blackhole delivery over active proofs by 4.67 percentage points in the prespecified thirty-fixture benchmark. Matched controls establish neither a unique censoring effect nor neural-policy superiority. The separate 180-run authenticated transfer study likewise shows no established attention advantage, while the 240-run factorial extension identifies a proof-traffic penalty under always-on listening. These findings apply to the declared feedback, traffic and radio models.',
 'Behavior change exposes a different boundary. In the 600-run Cooja delayed condition, active proofs deliver 79.11%, censored feedback 32.94% and naive feedback 54.89%; censored feedback with proof generation enabled delivers 32.83%. Censoring alone cannot reduce earlier positive trust when later observations disappear. Fixed-point MCU execution and fast private receipts improve narrower execution paths, but the original 90 ms signed-receipt target remains unmet. Physical RF, measured energy, production key storage and deployment performance remain unvalidated.']
 conclusion.text=parts[0];following=doc.add_paragraph(parts[1])
 if following._p.pPr is not None:following._p.remove(following._p.pPr)
 if conclusion._p.pPr is not None:following._p.insert(0,deepcopy(conclusion._p.pPr))
 conclusion._p.addnext(following._p);changes.append(dict(before=old,after='\n'.join(parts)))
 for prefix,oldpart,newpart in [('Earlier architecture comparisons','A subsequent 600-run telemetry study','The earlier 600-run telemetry study'),('The benchmark supports replacing','The benchmark supports replacing','The original 750-run benchmark supports replacing')]:
  p=next(p for p in doc.paragraphs if p.text.startswith(prefix));old=p.text;p.text=p.text.replace(oldpart,newpart);changes.append(dict(before=old,after=p.text))
 anchor=next(p for p in doc.paragraphs if p.text=='6. Discussion');body=next(p for p in doc.paragraphs if p.text.startswith('The simulated network has'))._p.pPr
 def para(text,style=None):
  p=doc.add_paragraph(text,style=style)
  if style is None:
   if p._p.pPr is not None:p._p.remove(p._p.pPr)
   p._p.insert(0,deepcopy(body))
  anchor._p.addprevious(p._p)
  if style or text.startswith('Table '):p.paragraph_format.keep_with_next=True
  return p
 para('5.9 Cooja feedback and policy transfer','Heading 2')
 para('A separate Cooja evaluation [25] runs the frozen Q22 actor and weighted feedback ledger inside four native Contiki-NG motes. Thirty new simulation seeds cross five arms, benign or post-warm-up blackhole behavior, and native or scheduled processing time: 600 runs and 36,000 offered packets. The source offers sixty packets at 300 ms intervals; two feasible relays connect it to the sink. All arms share geometry, traffic, authenticated beacons, nominal energy/queue inputs and the same checkpoint. Censored-plus-proof retains the censored evidence rule while enabling the active arm’s proof-generation rule; realized traffic can still differ as routes change. This is a four-node policy-transfer experiment, not a repetition of the original network-scale benchmark.')
 para('Actual CSMA callbacks supply first-hop MAC outcomes. A radio-driver hook observes successfully received forwarding frames before destination filtering, rather than reading simulator-held forwarding histories. Passive evidence has strength 0.5 and a signed sink receipt strength 1; the Python ledger independently replays ACK gating, 150 ms expiry, maximum-strength deduplication and late upgrades. Active proofs are authenticated hop messages emitted after successful downstream MAC outcomes. Public fixture keys support trace verification, not deployment secrecy. Sixteen local pending routes expire after two seconds; closed feedback records can remain eligible for late upgrades until their slots are reused.')
 para('Native execution does not account for Cortex-M3 instruction time. The delayed setting schedules outbound data after 14 ms, other outbound messages after 1 ms, received public receipts after 81 ms, sink data processing after 31 ms and other received frames after 1 ms. These per-packet assumptions do not model serial CPU service, overload or measured compute energy. Table 11 reports means over thirty seeds; Fig. 12 shows paired delivery contrasts, with a shared horizontal scale within each attack setting. Pointwise bootstrap intervals use 20,000 seed resamples. Two-sided paired sign-flip tests use 100,000 draws, with Holm adjustment across sixteen contrasts within each endpoint family. Delivery is primary; critical sink-deadline failures and total transmit-state duration form separate secondary families. Four direct matched-proof contrasts per endpoint were added after the protocol freeze as exploratory checks, with a separate Holm family; they do not alter primary inference.')
 para('Table 11. Cooja transfer results.')
 t=doc.add_table(rows=1,cols=6);t.autofit=False;headers=['Setting','Feedback arm','Delivery base (%)','Delivery delayed (%)','Critical failure delayed (%)','TX delayed (s)'];widths=[2,3.2,2.6,2.6,3,2]
 for index,row in enumerate([headers]+table_rows(a)):
  cells=t.rows[0].cells if index==0 else t.add_row().cells
  for c,text,width in zip(cells,row,widths):
   c.width=Cm(width);c.text=text
   for p in c.paragraphs:
    p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1
    for run in p.runs:run.font.name='Times New Roman';run.font.size=Pt(8.5);run.bold=index==0
  pr=t.rows[index]._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if index==0:pr.append(OxmlElement('w:tblHeader'))
 borders=OxmlElement('w:tblBorders')
 for side in ['top','bottom','insideH','left','right','insideV']:
  x=OxmlElement('w:'+side);x.set(qn('w:val'),'single' if side in ['top','bottom','insideH'] else 'nil');x.set(qn('w:sz'),'4');borders.append(x)
 t._tbl.tblPr.append(borders);anchor._p.addprevious(t._tbl)
 para('Delivery counts unique sink packets; critical failure includes missing critical packets or the sink application delivery event later than 90 ms. TX is the sum of modeled transmit-state seconds over all four motes, not energy. Base uses no added instruction-cost delay; the delayed column uses the declared processing assumptions. Full per-arm intervals, proof counts and adjusted contrasts are retained with the raw traces.')
 active=next(x for x in a['contrasts'] if (x['endpoint'],x['attack'],x['timing'],x['arm'])==('delivery',1,1,'active'))
 proof=next(x for x in a['contrasts'] if (x['endpoint'],x['attack'],x['timing'],x['arm'])==('delivery',1,1,'censored_proof'))
 para('With scheduled processing and a blackhole relay, active proofs change delivery relative to censored feedback by '+effect(active)+'. Enabling proofs while retaining censored evidence changes delivery by '+effect(proof)+'. These are conditional transfer results; they neither replace the original Holm family nor establish a new neural-policy advantage. Across all captures, independent reconstruction differs by at most '+f"{a['max_feature_error']:.3g}"+' per feature, with maximum logit discrepancy '+f"{a['max_logit_error']:.3g}"+'.')
 p=para('');p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(out/'cooja_figures/fig12_cooja_paired_delivery.png'),width=Cm(16))
 para('Fig. 12. Paired Cooja delivery effects.')
 scope=para('The four-node geometry, one source, fixed nominal energy/queue/LQI telemetry, public test keys and unicast beacon roster limit generalization. The fixed-point port quantizes float timestamps to integer milliseconds; a beacon can be excluded about 1 ms earlier than its ideal acquisition timestamp. The corrected auditor reproduces that representation without changing captures or outcomes. Energest state durations do not establish current or battery consumption, and MAC ACKs are not newly cryptographically authenticated. Censoring preserves earlier positive evidence when later observations are absent; recovery from a previously cooperative relay that becomes silent therefore depends on other routing inputs or additional evidence.')
 p._p.addprevious(scope._p)
 ref=doc.add_paragraph('[25] F. Österlind, A. Dunkels, J. Eriksson, N. Finne, and T. Voigt, “Cross-level sensor network simulation with COOJA,” in Proceedings of the 31st IEEE Conference on Local Computer Networks, pp. 641–648, 2006. DOI: 10.1109/LCN.2006.322172.')
 template=next(p for p in doc.paragraphs if p.text.startswith('[24]'))._p.pPr
 if ref._p.pPr is not None:ref._p.remove(ref._p.pPr)
 ref._p.insert(0,deepcopy(template))
 (out/'cooja_integration_edits.json').write_text(json.dumps(changes,indent=2,ensure_ascii=False)+'\n')
 return dict(abstract_words=len(abstract.text.split()),figures=12,tables=11,references=25)
