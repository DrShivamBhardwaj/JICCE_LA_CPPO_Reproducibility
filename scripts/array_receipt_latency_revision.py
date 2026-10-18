"""Report measured emulation optimizations with separate authentication endpoints."""
from pathlib import Path
from copy import deepcopy
import json
from docx.shared import Cm,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def apply(doc,root):
 root=Path(root);changes=[];summary=json.loads((root/'outputs/cc2538_fast/receipt_profiles_summary.json').read_text());assert summary['status']=='passed_scoped_comparison'
 def replace(prefix,text=None,fn=None):
  p=next(p for p in doc.paragraphs if p.text.startswith(prefix));old=p.text;p.text=fn(old) if fn else text;changes.append(dict(before=old,after=p.text))
 replace('A separate four-node CC2538 pilot',text='The four-node CC2538 pilot executes observation construction, frozen attention inference, hop authentication and receipt handling inside the MCU. The host triggers traffic and beacon exchange and captures UART output. Earlier public fixtures deliver twenty packets and verified receipts, with fifty matching routing decisions. Table 10 compares the retained private-key FP32 streams with two later implementations; each profile offers one hundred packets in a three-hop chain and one hundred in a geometry with two feasible source choices. These are sequential functional diagnostics, not independent statistical replicates or a new confirmatory comparison; N is the offered count.')
 replace('Every MCU decision matches the frozen Python actor',text='Operation-level profiling identifies approximately 270 ms per-hop FP32 inference and 81 ms Ed25519 verification in the declared emulator. A separate Q22 implementation uses signed 32-bit weights and activations, 64-bit accumulators and an 8193-point tanh lookup table. Export checks bound arithmetic for input coordinates in [−32, 32]; nonfinite or out-of-domain inputs are rejected. The lookup table adds 32,772 bytes beyond the 176,644 parameter bytes. All five hundred retained decisions and the later MCU decisions match the FP32 reference at their recorded observations, with logit errors below 1e-4. This is tested approximation agreement, not universal action equivalence. Inference falls to approximately 12–13 ms with one candidate. Caching the provisioned sink signing key also removes repeated key derivation from the packet path.')
 replace('The first kernel evaluated masked candidates',text='The experimental fast profile adds a distinct 66-byte application frame: the fourteen-byte packet identity, a sixteen-byte source–sink HMAC, the twenty-byte hop header and its sixteen-byte hop tag. The sink emits this private receipt after verifying source data, then generates the original Ed25519 proof. Relays check the hop envelope and pending route and forward the private receipt opaquely; they assign no forwarding credit from it. Only the source authenticates the end-to-end MAC and settles its service debt early. The separate public proof remains subject to Ed25519 verification and the original MAC-ACK credit gate. A private MAC is not transferable public evidence, and its deadline endpoint is not the original signed-receipt endpoint.')
 replace('The pilot fixes source 0',text='Fixed-point inference reduces signed-receipt misses from 134/200 to 68/200; the 90 ms class still fails. In the fast profile, all two hundred private receipts meet their source deadlines, while all two hundred deferred Ed25519 proofs are eventually verified and sixty-eight miss the original signed deadline. Extra private receipts add three radio frames per chain packet or two per choice packet; no energy saving is inferred. Tested-path stack watermarks reach 6624 bytes, with 3968 static RAM bytes and intact guards under the 32 KiB linker budget. These are observed paths, not a worst-case stack bound.')
 # Insert the remaining scope paragraph immediately after the preceding one.
 anchor=next(p for p in doc.paragraphs if p.text.startswith('Fixed-point inference reduces signed-receipt misses'))
 text='Fresh pairwise keys and a signing seed are provisioned once over the trusted emulator UART; persisted host allocation reserves distinct boot sessions. The software rejects stale sessions, repeated provisioning, mismatched signing keys and forged fast MACs, and uses explicit wipes for temporary key material and UART buffers. Source replay state rolls across longer streams. None of this establishes protected MCU storage, authenticated physical provisioning or reboot-persistent device state. The pilot still fixes four nodes, nominal energy/queue inputs and test-triggered beacons. Physical RF/energy, overload, autonomous rendezvous and deployment behavior remain unvalidated; retrospective private hop-tag audits require the locally retained fixtures.'
 p=doc.add_paragraph(text)
 if p._p.pPr is not None:p._p.remove(p._p.pPr)
 if anchor._p.pPr is not None:p._p.insert(0,deepcopy(anchor._p.pPr))
 anchor._p.addnext(p._p)
 abstract=doc.paragraphs[next(i for i,p in enumerate(doc.paragraphs) if p.text=='Abstract')+1];old=abstract.text
 abstract.text=old.replace('A four-node MCU emulation pilot executes live routing and authentication but misses the two tighter receipt deadlines.','A fixed-point MCU actor reduces signed-receipt deadline misses; a separate source-authenticated receipt profile meets its tested deadlines while retaining deferred public signatures.')
 changes.append(dict(before=old,after=abstract.text));assert len(abstract.text.split())<=250
 replace('ACK-gated censored passive feedback passed the prespecified',fn=lambda s:s.replace('Successful delivery does not establish real-time service: receipt delays exceed the two tighter class budgets.','Fixed-point inference improves emulated latency, but signed receipts still miss the 90 ms target. A separate source-only MAC receipt meets the tested source deadlines while retaining deferred public proofs; the two endpoints provide different verification properties.'))
 replace('Table 10. Live MCU pilot results.',text='Table 10. MCU receipt profiles.')
 old_table=doc.tables[9];t=doc.add_table(rows=1,cols=7);t.autofit=False
 headers=['Profile / topology','N','Delivered / Ed proofs','Ed delay (ms)','Ed misses','Fast delay (ms)','Fast misses'];widths=[2.6,.9,2.5,2.4,1.6,2.4,1.6]
 rows=[]
 for x in summary['rows']:
  profile={'cc2538_v2':'FP32','cc2538_q22':'Q22','cc2538_fast':'Fast'}[x['profile']]
  rows.append([profile+' / '+x['topology'],str(x['offered']),f"{x['delivered']} / {x['ed25519_receipts']}",f"{x['signed_delay_ms'][0]}–{x['signed_delay_ms'][1]}",str(x['signed_deadline_misses']),f"{x['fast_delay_ms'][0]}–{x['fast_delay_ms'][1]}" if 'fast_delay_ms' in x else '—',str(x['fast_deadline_misses']) if 'fast_deadline_misses' in x else '—'])
 for col,w in zip(t.columns,widths):col.width=Cm(w)
 for index,row in enumerate([headers]+rows):
  cells=t.rows[0].cells if index==0 else t.add_row().cells
  for c,v,w in zip(cells,row,widths):
   c.width=Cm(w);c.text=v
   for p in c.paragraphs:
    p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1
    for run in p.runs:run.font.name='Times New Roman';run.font.size=Pt(8.5);run.bold=index==0
 for i,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if i==0:pr.append(OxmlElement('w:tblHeader'))
  if i<len(t.rows)-1:
   for c in row.cells:
    for p in c.paragraphs:p.paragraph_format.keep_with_next=True
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','bottom','insideH','left','right','insideV']:
  e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single' if edge in ['top','bottom','insideH'] else 'nil');e.set(qn('w:sz'),'4');borders.append(e)
 t._tbl.tblPr.append(borders);old_table._tbl.addprevious(t._tbl);old_table._tbl.getparent().remove(old_table._tbl)
 (root/'publication/array/receipt_latency_edits.json').write_text(json.dumps(changes,indent=2,ensure_ascii=False)+'\n')
 return dict(abstract_words=len(abstract.text.split()),figures=11,tables=10,references=24)
