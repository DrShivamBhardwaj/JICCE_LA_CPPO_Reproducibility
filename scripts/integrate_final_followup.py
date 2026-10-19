"""Integrate audited compiler/radio diagnostics and author-supplied declarations."""
from pathlib import Path
from io import BytesIO
from lxml import etree
from copy import deepcopy
import subprocess,json,hashlib,shutil,zipfile,re
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parents[1];P=R/'publication/array';B=R/'build/followup_integration';B.mkdir(parents=True,exist_ok=True)
BASE='c65cab682aef45c0a680e33716fa82ea2bdd48eb'
raw=subprocess.check_output(['git','show',BASE+':publication/array/Array_Manuscript.docx'],cwd=R)
assert hashlib.sha256(raw).hexdigest()=='0b9565876c54f65633782e56e444b00fc9e602fd38349965a4523a1c505b931a'
(B/'baseline.docx').write_bytes(raw);doc=Document(BytesIO(raw))
old_math=[etree.tostring(x) for x in doc.element.xpath('.//m:oMath')];old_tables=[t._tbl.xml for t in doc.tables]
a=json.loads((R/'outputs/mcu_deadline_v2/analysis.json').read_text());e=json.loads((R/'outputs/cooja_radio_energy_v1/analysis.json').read_text())
assert len(a['rows'])==2 and all(x['delivered']==100 and x['signed_receipts']==100 and x['signed_deadline_misses']==34 for x in a['rows'])
assert e['status']=='passed_600_capture_and_radio_state_rechecks' and e['runs']==600 and e['nodes_verified']==2400 and e['new_simulations_executed']==0
replacements=[]
for p in doc.paragraphs:
 before=p.text
 if 'Compiler-only O3 MCU execution still misses' in p.text:p.text=p.text.replace('Compiler-only O3 MCU execution still misses 68 of 200 original signed-receipt deadlines.','The O3 and LTO compiler profiles each miss 68 of 200 original signed-receipt deadlines.')
 if 'A compiler-only O3 pilot retains' in p.text:p.text=p.text.replace('A compiler-only O3 pilot retains 68/200 signed-receipt deadline misses.','Both O3 and LTO compiler profiles retain 68/200 signed-receipt deadline misses. The radio-current sensitivity in Section 5.12 also shows why lower transmit airtime alone cannot establish measured energy savings.')
 if before!=p.text:replacements.append({'before':before,'after':p.text})
anchor=next(p for p in doc.paragraphs if p.text=='6. Discussion')
def add(text='',style='Normal'):
 p=anchor.insert_paragraph_before(text,style)
 if style=='Normal':
  p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(6);p.paragraph_format.space_before=Pt(0)
 return p
add('5.12 Compiler optimization and radio current sensitivity','Heading 2')
add('A further compiler diagnostic was frozen at commit b52842496dc625ad30ff95bdefad4c107366080e before execution in workflow 39179074994. Link-time optimization (LTO) and default loop optimization supplement O3; the Q22 actor, Ed25519 endpoint, clock model, traffic and 90/350/1500 ms deadlines remain unchanged. All 200 offered packets produce source-verified signed receipts, all 500 routing decisions match the frozen reference, and seven authenticated negative-control categories pass in each topology. Table 13 compares the retained diagnostics. These sequential runs use fresh private keys and are not independent statistical replicates of a compiler speedup.')
cap=add('Table 13. Signed receipt compiler diagnostics.');cap.paragraph_format.keep_with_next=True
rows=[['Profile','Topology','Signed delay (ms)','Misses / offered','Decisions'],['O3','Chain','331–344','34 / 100','300'],['O3','Choice','239–247','34 / 100','200'],['O3 + LTO','Chain','299–308','34 / 100','300'],['O3 + LTO','Choice','213–220','34 / 100','200']]
t=doc.add_table(rows=0,cols=5);t.style=doc.tables[0].style
for row in rows:
 cells=t.add_row().cells
 for cell,text in zip(cells,row):
  cell.text=text
  for p in cell.paragraphs:
   p.paragraph_format.keep_with_next=True
   for run in p.runs:run.font.size=Pt(9)
 for cell in cells:
  pr=cell._tc.get_or_add_tcPr()
 for r in t.rows:
  pr=r._tr.get_or_add_trPr()
  if not pr.find(qn('w:cantSplit')) is not None:pr.append(OxmlElement('w:cantSplit'))
for c in t.rows[0].cells:
 for r in c.paragraphs[0].runs:r.bold=True
borders=OxmlElement('w:tblBorders')
for side in ['top','bottom','insideH']:
 edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');borders.append(edge)
t._tbl.tblPr.append(borders)
anchor._p.addprevious(t._tbl)
add('The LTO profile still misses every critical-class deadline: 34/100 packets in each topology. Packet-matched sink signing plus source signature verification occupies at least 104 ms in the chain and 105 ms in the choice geometry; both medians are 106 ms. These two serial dependencies alone exceed 90 ms before relay verification, routing and radio overhead are added. The observations describe this emulator and implementation; they are not universal lower bounds on Ed25519 or physical-device timing. Independent public-key checks verify all 200 sink signatures, match all 200 accepted source receipts to received signed frames, and reject 200 one-bit signature mutations. Private hop tags cannot be retrospectively reverified without the intentionally unexported per-run secrets.')
add('A separate exploratory reanalysis checks all 600 retained historical Cooja runs: thirty seeds per attack, timing and feedback-arm stratum, 2,400 final node records and 36,000 offered packets. Every decompressed capture matches its archived hash; transmit/listen state accounting and source/sink endpoint counts match the retained audits. No new simulations are represented by this reanalysis. Fig. 15 relates descriptive delivery means to summed node transmit time. The original active and censored arms retain their historical semantic differences; this plot is not a substitute for the matched-factorial inference in Section 5.11.')
folder=P/'followup_figures';folder.mkdir(exist_ok=True)
for num,name,caption,width in [(15,'radio_delivery','Fig. 15. Delivery and transmit airtime.',6.2),(16,'assumed_current_sensitivity','Fig. 16. Assumed radio current sensitivity.',6.2)]:
 for ext in ['png','pdf','svg']:shutil.copyfile(R/f'outputs/cooja_radio_energy_v1/{name}.{ext}',folder/f'fig{num}_{name}.{ext}')
 if num==16:
  add('For a common supply voltage V (volts) and constant listen current I_L (amperes) across nodes, the radio-only model is E_radio = V I_L (T_L + r T_TX), where T_L and T_TX are summed node state durations in seconds and r = I_TX/I_L is dimensionless. Fig. 16 varies the assumed ratio over 0.5, 1, 2 and 3. The plotted relative values require no assumed absolute current or voltage. At r = 1, every arm has equal radio energy over the always-on accounting window; rankings can reverse across this point. CPU, cryptography, sensing, regulator and host consumption are excluded. These are conditional model calculations, not measured currents, joules or battery lifetimes. Coincident plot markers indicate equal means.')
 p=add();p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(folder/f'fig{num}_{name}.png'),width=Inches(width));add(caption)
add('The retained datasets therefore support virtual radio-state and endpoint checks. Calibrated physical RF/current acquisition, board-specific protected keys, authenticated provisioning, firmware authenticity and reset-safe persistent MCU sessions remain unvalidated. The native Cooja timing assumptions do not supply a serial CPU model or physical calibration.')
ack=next(p for p in doc.paragraphs if p.text=='Acknowledgments')
ack.insert_paragraph_before('Author contributions','Heading 1')
ack.insert_paragraph_before('Shivam Bhardwaj, Abhishek Kumar Pandey and Abhinandan Tripathi contributed equally to this work.','Normal')
for paragraph in doc.paragraphs:
 if paragraph.text.startswith('For a common supply voltage'):
  content=paragraph.text;paragraph.clear()
  for part in re.split(r'(E_radio|I_L|T_L|T_TX|I_TX)',content):
   if part in ['E_radio','I_L','T_L','T_TX','I_TX']:
    base,sub=part.split('_');paragraph.add_run(base);run=paragraph.add_run(sub);run.font.subscript=True
   else:paragraph.add_run(part)
assert old_math==[etree.tostring(x) for x in doc.element.xpath('.//m:oMath')]
assert old_tables==[t._tbl.xml for t in doc.tables[:12]]
assert len(doc.inline_shapes)==16 and len(doc.tables)==13
doc.save(P/'Array_Manuscript.docx')
(P/'followup_edits.json').write_text(json.dumps(replacements,indent=2)+'\n')
(P/'Array_Manuscript.md').write_text('\n\n'.join(p.text for p in doc.paragraphs)+'\n')
with zipfile.ZipFile(P/'Array_Figures_PNG_PDF_SVG.zip','w',zipfile.ZIP_DEFLATED) as z:
 for name in ['figures','review_figures','postfreeze_figures','cooja_figures','remediation_figures','followup_figures']:
  for f in sorted((P/name).glob('fig*')):
   if f.suffix in ['.png','.pdf','.svg']:z.write(f,str(f.relative_to(P)))
(P/'followup_integration_audit.json').write_text(json.dumps(dict(status='passed_content_preservation_checks',baseline_commit=BASE,baseline_docx_sha256=hashlib.sha256(raw).hexdigest(),docx_sha256=hashlib.sha256((P/'Array_Manuscript.docx').read_bytes()).hexdigest(),native_math_objects_preserved=len(old_math),old_tables_preserved=12,figures=16,tables=13,lto_misses=68,radio_runs_reanalyzed=600,physical_validation=False,equal_contribution_source='Corresponding-author statement 2026-10-19T06:35:35+05:30',competing_interest_statement='No competing interests; supplied by corresponding author'),indent=2)+'\n')
print('Integrated:',len(doc.inline_shapes),'figures,',len(doc.tables),'tables; preserved',len(old_math),'native math objects')
