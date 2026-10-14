from pathlib import Path
import json
import re
import sys
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'publication'
OUT.mkdir(exist_ok=True)
DATA = ROOT / 'outputs/jicce_validation'
a = pd.read_csv(DATA/'paired_analysis.csv')
s = pd.read_csv(DATA/'seed_summary.csv')
labels = {'attention_minus_capacity_mlp':'Capacity-matched MLP',
          'attention_minus_attention_no_risk':'Attention without risk',
          'attention_minus_attention_feedback80':'Attention with 80% feedback'}
plt.rcParams.update({'font.family':'DejaVu Serif','font.size':12,'axes.spines.top':False,'axes.spines.right':False})
fig, axs = plt.subplots(1,2,figsize=(10,3.5),layout='constrained')
for ax, metric, name in zip(axs,['class_0_mean_delay_ms','class_0_delay_violation_rate'],['Mean delivered delay difference (ms)','Total violation difference (percentage points)']):
    d=a[(a.setting=='nominal')&(a.metric==metric)].copy()
    scale=100 if metric.endswith('rate') else 1
    for i, row in enumerate(d.itertuples()):
        ax.errorbar(row.mean_difference*scale,i,xerr=[[max(0,(row.mean_difference-row.ci95_low)*scale)],[max(0,(row.ci95_high-row.mean_difference)*scale)]],fmt='o',color='#006d77',capsize=4)
    ax.set_yticks(range(len(d)),[labels[x] for x in d.comparison] if ax is axs[0] else ['']*len(d))
    ax.axvline(0,color='gray',ls='--',lw=1);ax.set_xlabel(name);ax.grid(axis='x',alpha=.2)
for ext in ['svg','pdf','eps','png']:fig.savefig(OUT/'figures'/f'fig6.{ext}',dpi=900 if ext=='png' else 300)
plt.close(fig)

def format_doc(d):
    for sec in d.sections:
        sec.page_width=Cm(21);sec.page_height=Cm(29.7)
        sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(2.5)
        cols=sec._sectPr.find(qn('w:cols'))
        if cols is not None:cols.set(qn('w:num'),'1')
    for node in list(d.element.xpath('.//w:pBdr')):
        node.getparent().remove(node)
    for sty in d.styles:
        for node in list(sty.element.xpath('.//w:pBdr')):
            node.getparent().remove(node)
    for sty in ['Normal','Body Text']:
        d.styles[sty].font.name='Times New Roman';d.styles[sty].font.size=Pt(9.5)
        d.styles[sty].paragraph_format.line_spacing=2
    for sty in ['Title','Heading 1','Heading 2','Heading 3']:
        d.styles[sty].font.color.rgb=RGBColor(0,0,0)
    for p in d.paragraphs:
        p.paragraph_format.space_after=Pt(4)
        if p.text and not p.text.startswith('Fig.'):
            p.paragraph_format.line_spacing=2
        for run in p.runs:
            run.font.name='Times New Roman'
            run.font.color.rgb=RGBColor(0,0,0)
            if p.style.name not in ['Title','Heading 1','Heading 2','Heading 3']:run.font.size=Pt(9.5)
    for shape in d.inline_shapes:
        if shape.width > Cm(16):
            shape.height=int(shape.height*Cm(16)/shape.width);shape.width=Cm(16)
    for t in d.tables:
        pr=t._tbl.tblPr
        for old in list(pr.findall(qn('w:tblBorders'))):pr.remove(old)
        borders=OxmlElement('w:tblBorders')
        for edge in ['top','bottom','left','right','insideH','insideV']:
            e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single' if edge in ['top','bottom','insideH'] else 'nil');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'888888');borders.append(e)
        pr.append(borders)
        if t.rows:
            trpr=t.rows[0]._tr.get_or_add_trPr()
            if trpr.find(qn('w:tblHeader')) is None:trpr.append(OxmlElement('w:tblHeader'))
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing=1
                    for run in p.runs:run.font.name='Times New Roman';run.font.size=Pt(8)
    for sec in d.sections:
        footer=sec.footer.paragraphs[0]
        for child in list(footer._p):footer._p.remove(child)
        footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
        run=footer.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');run._r.addnext(fld)

source=Path(sys.argv[1])
d=Document(source)
d.paragraphs[0].text='Local Attention Routing for Wireless Sensor Networks with Reliability and Latency Constraints'
d.paragraphs[0].style=d.styles['Title']
capacity=a[(a.setting=='nominal')&(a.comparison=='attention_minus_capacity_mlp')]
delay=capacity[capacity.metric=='class_0_mean_delay_ms'].iloc[0]
violation=capacity[capacity.metric=='class_0_delay_violation_rate'].iloc[0]
d.paragraphs[6].text=(
    'Wireless sensor routing must balance packet delivery, delay and energy under uncertain forwarding behavior. '
    'We evaluate a local-attention constrained proximal policy optimization router using class-specific service debt and online Bayesian forwarding risk. '
    'Its shared candidate encoder gives an equivariant actor and invariant critic. In the historical 30-seed comparison with a same-width multilayer perceptron, '
    'critical-class total violations fell by 3.125 percentage points, while mean delivered delay increased by 1.016 ms; the joint superiority criterion failed. '
    f'A separate 30-seed validation compared a model with approximately matched parameter count and tested risk removal, incomplete feedback and frozen-policy transfer. '
    f'Against the capacity comparator, the delay difference was {delay.mean_difference:+.3f} ms and the total-violation difference was {100*violation.mean_difference:+.3f} percentage points. '
    'Neither capacity-comparison endpoint was significant after the planned multiplicity correction. Risk removal degraded reliability. These findings concern simulated routing and establish neither live intrusion detection nor guaranteed service constraints.')
d.paragraphs[7].text='Index Terms: Constrained reinforcement learning, Local attention, Risk-aware routing, Wireless sensor networks'
d.paragraphs[5].insert_paragraph_before('Running title: Risk aware local attention routing')

# Insert a separate validation section before the historical discussion.
anchor=next(p for p in d.paragraphs if p.text=='9. Discussion')
new=[]
def para(text,style=None):
    p=d.add_paragraph(text,style);new.append(p._p);return p
para('9. Independent reviewer-led validation','Heading 1')
para('This validation was specified separately after the historical AF-v3 and AF-v4 programs. The original gates and outcomes remain unchanged. Thirty new evaluation seeds, 41001–41030, were fixed before the new outcomes were generated. Each arm received eight updates of 50 training episodes and three PPO epochs per update. Trained weights were evaluated for 240 episodes in each setting and retained as checkpoints.')
counts=s[s.setting=='nominal'].groupby('policy').model_parameter_count.first()
para(f'The attention model contains {int(counts.attention):,} parameters. The nearest-size fixed-slot MLP contains {int(counts.capacity_mlp):,}. Its width was chosen from parameter count alone. This controls total parameter count approximately; compute, layer structure and inductive bias remain different. Two additional attention arms remove risk from candidate features, preselection and reward, or observe only 80% of forwarding outcomes. The latter is a missing-feedback sensitivity analysis, not a live-traffic implementation.')
para('Separate seeded generators govern traffic selection, queue dynamics, attack transitions, radio draws and feedback omission. Radio consumption therefore does not alter later traffic-class draws. Policy-dependent radio and attack draws remain possible, and source eligibility can diverge after node death. The experiment does not replay an identical per-edge external-event schedule.')
para('The nominal comparison family consists of the six paired tests formed by three comparators and two critical-class endpoints. Holm correction is applied jointly to these six tests. Paired 95% intervals are unadjusted. Transfer evaluations retain the trained weights and change only the node count to 200 in the same area or the malicious-node ratio to 10% or 40%. These secondary results measure sensitivity; increasing density is not a density-preserving scalability test.')
para(f'Against the capacity-matched MLP, attention changed mean delivered delay by {delay.mean_difference:+.4f} ms (95% CI {delay.ci95_low:+.4f} to {delay.ci95_high:+.4f}; Holm p={delay.holm_primary_p:.4g}). Total violations changed by {100*violation.mean_difference:+.4f} percentage points (95% CI {100*violation.ci95_low:+.4f} to {100*violation.ci95_high:+.4f}; Holm p={violation.holm_primary_p:.4g}). Table 8 reports all nominal comparisons, and Fig. 6 shows their effect intervals. Unfavorable and inconclusive contrasts are retained.')
para('Table 8. New nominal validation effects for attention minus each comparator.')
t=d.add_table(rows=1,cols=5);new.append(t._tbl)
for c,v in zip(t.rows[0].cells,['Comparator','Endpoint','Difference','95% CI','Holm p']):c.text=v
for row in a[a.setting=='nominal'].query("metric in ['class_0_mean_delay_ms','class_0_delay_violation_rate']").itertuples():
    scale=100 if row.metric.endswith('rate') else 1
    vals=[labels[row.comparison],'Delay ms' if scale==1 else 'Violation pp',f'{row.mean_difference*scale:+.3f}',f'[{row.ci95_low*scale:+.3f}, {row.ci95_high*scale:+.3f}]',f'{row.holm_primary_p:.4g}']
    for c,v in zip(t.add_row().cells,vals):c.text=v
p=d.add_paragraph();p.add_run().add_picture(str(OUT/'figures/fig6.png'),width=Cm(16));new.append(p._p)
para('Fig. 6. New nominal paired effects with unadjusted 95% intervals across 30 seeds. Negative differences favor attention.')
para('Table 9. Frozen-policy transfer summaries across 30 seeds.')
t=d.add_table(rows=1,cols=5);new.append(t._tbl)
for c,v in zip(t.rows[0].cells,['Setting','Policy','Delivery %','C0 delay ms','C0 violation %']):c.text=v
for (setting,policy), row in s[s.policy.isin(['attention','capacity_mlp'])].groupby(['setting','policy']).mean(numeric_only=True).iterrows():
    vals=[setting,policy,f'{row.delivery_rate*100:.2f}',f'{row.class_0_mean_delay_ms:.2f}',f'{row.class_0_delay_violation_rate*100:.2f}']
    for c,v in zip(t.add_row().cells,vals):c.text=v
para('Forwarding observations still come directly from simulator outcomes and expected link reliability. Their collection, authentication and communication cost are not measured. The available evidence therefore supports an online simulated feedback loop, with missing-feedback sensitivity, rather than practical validation of an intrusion-detection system.')
for el in new:anchor._p.addprevious(el)
for p in d.paragraphs:
    if re.match(r'^(9|10|11|12)\.',p.text) and p._p not in new:
        m=re.match(r'^(\d+)(.*)',p.text);p.text=str(int(m.group(1))+1)+m.group(2)
    if p.text.startswith('Network-size transfer and deployment-overhead experiments were likewise conditional'):
        p.text='Historical AF-v3 transfer and overhead experiments were not run after its failed gate. The separate new validation reports limited frozen-policy sensitivity to node count and attack prevalence. No hardware benchmark or complete deployment-overhead evaluation has been performed.'
    if p.text.startswith('The model was not parameter-count matched:') or p.text.startswith('The models are state-matched but not parameter-matched:'):
        p.text='The historical comparison was not parameter-count matched: attention had 63,368 parameters and the same-width MLP had 19,791. The separate new validation approximately matches total parameter count and is reported in Section 9. Neither design isolates architecture at equal computation cost.'
    if p.text.startswith('The MLP is matched on state content and training budget, not on trainable parameter count.'):
        p.text='The historical MLP is matched on state content and training budget rather than parameter count. The new comparison controls total capacity approximately; neither comparison establishes an isolated effect of representation at equal computation cost.'
    if p.text.startswith('The primary threat model fixes a 25% mixed-attack setting.'):
        p.text='Historical confirmation uses a 25% mixed-attack setting. The separate validation changes attack prevalence to 10% and 40% but retains the same attack mechanisms. Robustness to unseen attack types, collusion, adaptive evasion, Sybil attacks and wormholes remains unknown.'
    if p.text.startswith('Data and code availability:'):
        p.text='Data and code availability: Source code, configurations, historical AF-v3 raw outputs, new validation trajectories, seed summaries, trained checkpoints and analysis scripts are available at https://github.com/DrShivamBhardwaj/JICCE_LA_CPPO_Reproducibility. Historical and new evidence are separated, and the original failed joint criterion is retained.'
    if p.text.startswith('The accompanying package provides the AF-v3 and AF-v4 source snapshots'):
        p.text='The public package contains the current source, configurations, historical AF-v3 outputs, new validation data and checkpoints, analysis scripts and a file checksum manifest. Historical source commits and AF-v4 records are identified by links to their original provenance. Exact historical training requires the original source commit and environment; the current validation is not a bitwise reproduction claim.'
    if p.text.startswith('The authors used  OpenAI'):
        p.text='The authors used OpenAI tools to assist with language editing and preparation of code and analysis. They reviewed the methods, experimental outputs and manuscript and take responsibility for the final content.'
    if p.text.startswith('LA-CPPO combines a candidate-equivariant local actor'):
        p.text=(f'LA-CPPO combines a candidate-equivariant actor, invariant critics, class-specific service debt and forwarding-risk estimates. The historical same-width comparison showed a reliability benefit but failed the joint superiority criterion. In the separate capacity-controlled validation, mean delivered delay changed by {delay.mean_difference:+.3f} ms and total violations by {100*violation.mean_difference:+.3f} percentage points. Neither capacity-comparison co-primary endpoint was significant after Holm correction. Risk removal substantially reduced reliability, while incomplete-feedback differences were inconclusive. Transfer outcomes remain setting-specific. Neither study establishes live IDS performance, absolute service-budget satisfaction or deployment guarantees.')
refs=next(p for p in d.paragraphs if p.text=='References')
appendix=[]
for text in ['Appendix A Historical development and constraint rates',
             'The AF-v4 development output retained two unsuccessful candidates. Their mean-delay differences against the slack-aware MLP were +3.225549 ms and +1.628427 ms. Neither passed the frozen development selection gate; no candidate was selected. Raw development comparisons, episode logs and selection_decision.json are supplied under data/historical_af_v4. These outcomes are not pooled with either 30-seed study.',
             'Table 10. Historical attention constraint rates across 30 seeds.']:
    p=d.add_paragraph(text);appendix.append(p._p)
t=d.add_table(rows=1,cols=4);appendix.append(t._tbl)
for c,v in zip(t.rows[0].cells,['Class','Total violation rate','Loss rate','Budget for each channel']):c.text=v
hist=pd.read_csv(ROOT/'data/historical_af_v3/seed_summary.csv')
hist=hist[hist.policy=='af_v3_attention_cppp']
for c in range(3):
    vals=[str(c),f'{hist[f"class_{c}_delay_violation_rate"].mean():.4f}',f'{hist[f"class_{c}_loss_rate"].mean():.4f}',str([.015,.04,.10][c])]
    for cell,v in zip(t.add_row().cells,vals):cell.text=v
for el in appendix:refs._p.addprevious(el)
for p in d.paragraphs:
    for run in p.runs:
        updated=run.text.replace('Supplementary Note S1','Appendix A').replace('Supplementary Table S2','Table 10').replace('Figure S2 summarizes this decomposition.','Table 7 summarizes this decomposition.')
        if updated!=run.text:run.text=updated
format_doc(d)
d.save(OUT/'JICCE_LA_CPPO_Revised_Manuscript.docx')

r=Document()
r.add_heading('JICCE resubmission evidence record',0)
r.add_paragraph('Related rejected manuscript JICCE-26-0057. This record accompanies a substantially redesigned study for consideration as a new submission. It does not claim that the original ROA–Maxout pipeline has been validated or that its small ROA effect has increased.')
for heading,text in [
('Reviewer A major comment 1', 'The old class-conditioned score replay has been removed from the redesigned router. The current code uses Bayesian risk updated from simulated forwarding outcomes and queue signals. New risk-removal and 80% feedback-observation experiments test the consequences of that feedback. They are simulated validation. Live telemetry extraction, distributed feedback authentication and measured feedback cost remain unresolved; the manuscript does not describe the method as a validated deep IDS.'),
('Reviewer A major comment 2', 'The redesigned study has no supervised node-ID-suffix dataset partition. Independent environment initialization separates training and evaluation. A separately frozen validation uses 30 new evaluation seeds, approximately matches comparator capacity and tests frozen weights on changed network sizes and attack prevalence. Separate random generators prevent radio draws from shifting traffic selection. Exact policy-independent external-event replay and broader field generalization remain outside the completed evidence.'),
('Reviewer A major comment 3', f'ROA is no longer the proposed routing method. Its earlier 0.074 percentage-point advantage is not reinterpreted as meaningful superiority. The new method is LA-CPPO, compared with a capacity-matched MLP. The new delay effect is {delay.mean_difference:+.4f} ms and the total-violation effect is {100*violation.mean_difference:+.4f} percentage points, with the intervals and Holm-adjusted tests in Table 8. The original AF-v3 failed joint criterion remains reported.'),
('Reviewer A minor comments', 'The revised initial-submission document uses A4 paper, 2.5 cm margins, 9.5-point text and double-spaced body paragraphs, following the author instructions distributed with the society’s available template bundle. The new statistical artwork is supplied in SVG, PDF and EPS as well as high-resolution PNG. The bundle is older; final compliance must be checked against the current submission portal. Historical figures have been preserved.'),
('Reviewer B', 'The prior response is retained as a historical record. Its acceptance of the earlier revisions is not treated as approval of the substantially redesigned model or the new experiments.'),
('Reproducibility and remaining limits', 'Repository: https://github.com/DrShivamBhardwaj/JICCE_LA_CPPO_Reproducibility. The package separates historical and new outputs, retains every planned seed/arm, includes 120 trained checkpoints and supplies scripts to regenerate numerical summaries. There are 42 passing implementation tests. Source commit and dependency versions are recorded. No hardware experiments, measured telemetry costs or actual intrusion-detection accuracy are claimed.')]:
    r.add_heading(heading,1);r.add_paragraph(text)
format_doc(r)
r.save(OUT/'JICCE_Resubmission_Evidence_Record.docx')
