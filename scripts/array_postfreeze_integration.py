"""Integrate the separately audited post-freeze experiments into the active draft."""
from pathlib import Path
from copy import deepcopy
import json
from docx.shared import Pt,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def apply(doc,root):
    root=Path(root);changes=[]
    def replace(prefix,new=None,fn=None):
        p=next(p for p in doc.paragraphs if p.text.startswith(prefix));old=p.text
        p.text=fn(old) if fn else new;changes.append(dict(before=old,after=p.text))
    abstract=doc.paragraphs[next(i for i,p in enumerate(doc.paragraphs) if p.text=="Abstract")+1]
    old=abstract.text;abstract.text=old.replace("An exploratory 60–300 s extension", "A further 240-run factorial study separates proof traffic from evidence handling; adding proofs reduces delivery by 6.11 percentage points with always-on listening. An isolated Cortex-M3 harness reproduces actor decisions and executes HMAC-SHA256 and Ed25519. An exploratory 60–300 s extension")
    changes.append(dict(before=old,after=abstract.text))
    replace("A four-second horizon, a central sink",fn=lambda s:s.replace("The experiments provide neither a long-horizon reclamation policy nor measured embedded inference cost.","The frozen experiments provide neither a long-horizon reclamation policy nor measured embedded inference cost. Sections 5.6–5.7 add bounded-state tests and isolated instruction-level execution; integrated firmware costs remain unmeasured."))
    replace("The inference actor contains 44,161",fn=lambda s:s.replace("target inference and measured energy require separate validation.","integrated target inference and measured energy require separate validation."))
    replace("ACK-gated censored passive feedback passed the prespecified",fn=lambda s:s.replace("embedded policy/crypto execution and measured energy remain unvalidated.","integrated on-board policy/crypto execution and measured energy remain unvalidated.")+" The 240-run factorial extension identifies a proof-traffic delivery penalty with always-on listening, while severe loss under periodic sleeping limits transfer of that finding. Isolated Cortex-M3 execution verifies actor outputs and cryptographic primitives, not a complete embedded radio system.")
    replace("The 150 ms evidence timeout",fn=lambda s:s.replace("device clock offset or drift, which remain untested.","physical clock offset or drift, which remain unmeasured; Section 5.7 tests a declared virtual envelope."))
    replace("A useful next step is to decide",new="Section 5.6 separates proof generation from the declared evidence rule, but its periodic sleep primitive lacks rendezvous. An observation policy could choose silence, passive listening or an authenticated probe according to the expected reduction in delivery loss relative to measured radio and processing cost. A discriminating deployment experiment would integrate rendezvous, test adversarial receipts and clock uncertainty, and reserve new physical-board settings. This would evaluate evidence acquisition under execution costs absent from the isolated harness.")
    replace("The action and ledger audit checked",fn=lambda s:s.replace("the expanded implementation now has 83 passing tests","the pre-extension implementation had 83 passing tests")+" The post-freeze suite in Sections 5.6–5.7 has 107 passing tests.")
    anchor=next(p for p in doc.paragraphs if p.text=="6. Discussion")
    body_template=next(p for p in doc.paragraphs if p.text.startswith("The simulated network has"))._p.pPr

    def para(text,style=None):
        p=doc.add_paragraph(text,style=style)
        if style is None:
            current=p._p.pPr
            if current is not None:p._p.remove(current)
            p._p.insert(0,deepcopy(body_template))
        previous=anchor._p.getprevious();anchor._p.addprevious(p._p)
        if previous is not None and previous.tag==qn("w:tbl"):p.paragraph_format.space_before=Pt(6)
        if text.startswith("Table "):p.paragraph_format.keep_with_next=True
        if style and style.startswith("Heading"):p.paragraph_format.keep_with_next=True
        return p
    def table(headers,rows,widths):
        t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
        for c,h,w in zip(t.rows[0].cells,headers,widths):c.text=h;c.width=Cm(w)
        for row in rows:
            cells=t.add_row().cells
            for c,v,w in zip(cells,row,widths):c.text=str(v);c.width=Cm(w)
        for i,row in enumerate(t.rows):
            trPr=row._tr.get_or_add_trPr();el=OxmlElement("w:cantSplit");trPr.append(el)
            if i==0:el=OxmlElement("w:tblHeader");trPr.append(el)
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1.
                    for run in p.runs:run.font.name="Times New Roman";run.font.size=Pt(9);run.bold=i==0
        pr=t._tbl.tblPr;b=OxmlElement("w:tblBorders")
        for edge in ["top","bottom","insideH"]:
            x=OxmlElement("w:"+edge);x.set(qn("w:val"),"single");x.set(qn("w:sz"),"4");x.set(qn("w:color"),"999999");b.append(x)
        for edge in ["left","right","insideV"]:
            x=OxmlElement("w:"+edge);x.set(qn("w:val"),"nil");b.append(x)
        pr.append(b);anchor._p.addprevious(t._tbl)
        return t
    factorial=json.loads((root/"outputs/validation_20261018/factorial_analysis.json").read_text())
    para("5.6 Factorial traffic controls and power gating","Heading 2")
    para("A separate ns-3.47 extension tests proof traffic while holding the evidence rule fixed within each contrast. Its amended source and 30-test analysis plan were committed as 752c11e7fa3c6af427ffe4edb8fa37d3fbcd5bc7 before fixtures 89001–89030 were evaluated. An initial proof-callback error stopped a preliminary 88001 batch; those partial captures and technical pilot 76901 are excluded. Thirty fixtures, four arms and two duty settings yield 240 four-second runs, each offering 60 packets. All arms use the same bounded positive-credit rule, ACK gate, 128-record global simulation-ledger capacity, 2 s retention and frozen attention actor. Unlike Eq. (4), each record clips q once to [0.05, 1] and weights each eligible uncredited increment in s by that fixed value. Proof frames consume traffic but never enter the ledger. Censored versus naive mode changes expiry penalties and late-evidence handling. Table 8 reports arm means; Fig. 11 shows the proof-traffic contrasts.")
    para("With always-on listening, adding proofs reduces censored-mode delivery by 6.11 percentage points (unadjusted paired 95% interval −8.94 to −3.28; Holm p=0.00349), and naive-mode delivery by 5.61 points (−7.94 to −3.29; Holm p=0.000852). No evidence-mode or interaction contrast survives the separate 30-test correction. These results identify a traffic contribution in this extension, not the complete cause of the original 4.67-point gain. The nominal 20% awake schedule rejects requests while asleep and has no rendezvous protocol; delivery falls to approximately 13–14%. Busy MAC transactions defer sleep, so nominal duty is not actual PHY residence.")
    para("Table 8. Factorial arm means.")
    rows=[]
    for x in factorial["arm_means"]:
        rows.append([f'{100*x["duty_fraction"]:.0f}%',"Censored" if x["arm"].startswith("censored") else "Naive","On" if x["arm"].endswith("with_proof") else "Off",f'{x["delivery_percent"]:.2f}',f'{x["critical_violation_percent"]:.2f}',f'{x["radio_energy_j"]:.6f}'])
    table(["Nominal awake","Evidence","Proofs","Delivery (%)","Critical failure (%)","Radio energy (J)"],rows,[2.0,2.3,1.5,2.5,3.0,3.0])
    p=para("");p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(root/"publication/array/postfreeze_figures/fig11_factorial_delivery.png"),width=Cm(15.6))
    para("Fig. 11. Proof-traffic delivery effects.")
    para("Hard power gating stops the experimental bridge at each sensor’s 0.5 J budget. Twenty-seven further 60 s runs cross three fixtures, three feedback arms and three nominal awake fractions. All remain within the 50 J aggregate sensor budget, with zero post-cut transmissions; the audit reconciles 2,727 radio-energy identities. The bounded ledger has zero final watches in all nine runs, whereas legacy controls retain records. All sensors still exhaust their budgets within 60 s. These tests address power accounting and ledger retirement, not sustained service or a whole-system memory bound.")
    para("5.7 Clock envelope and Cortex M3 execution","Heading 2")
    para("Optional freshness handling uses an externally justified pairwise clock bound. A 40 ms offset bound plus 100 ppm relative drift over 300 s gives 70 ms tolerance. All 675 tested offset/drift combinations accept authentic fresh beacons, while replay and forged-tag checks remain effective. The wider window can also admit older messages near the expiry boundary. It neither synchronizes clocks nor verifies their physical drift; the original default behavior remains unchanged.")
    para("A bare-metal C harness executes the 44,161-parameter actor, HMAC-SHA256 and Monocypher 4.0.3 Ed25519 on QEMU 11.1.2’s MPS2 AN385 Cortex-M3 model [23, 24]. Eight retained radio decisions and eight candidate-permutation controls produce 128 logits; all 16 actions match PyTorch and span eight candidate indices. The largest logit difference is 1.7881393432617188e-7, below the 1e-5 tolerance. HMAC matches RFC 4231 case 1; Ed25519 signatures match independent Python verification, and a corrupted signature is rejected. Table 9 records the execution checks.")
    para("Table 9. Isolated target checks.")
    table(["Check","Observed result"],[["Actor outputs","16/16 actions; 128 logits"],["Maximum logit error","1.7881393432617188e-7"],["HMAC / Ed25519","Known-answer / sign and verify passed"],["Text and read-only data","341,600 bytes"],["Observed stack watermark","3,080 bytes"],["Linker limits","512 KiB flash; 32 KiB SRAM"]],[6.4,8.0])
    para("The image omits radio buffers, RTOS state and integrated protocol memory; its observed stack watermark is not a worst-case bound. Public test keys are used only for validation. This isolated instruction harness and the CC2538 UART/radio bridge in Section 5.4 are distinct experiments. Neither their separate success nor the activation approximation’s agreement on these vectors validates a complete embedded routing system, physical timing or measured energy.")
    ref_template=next(p for p in doc.paragraphs if p.text.startswith("[22]"))._p.pPr
    p=doc.add_paragraph('[23] L. Vaillant et al., “Monocypher,” release 4.0.3, 2026. Available: https://monocypher.org/changelog. Accessed October 18, 2026.')
    if p._p.pPr is not None:p._p.remove(p._p.pPr)
    p._p.insert(0,deepcopy(ref_template))
    p=doc.add_paragraph('[24] QEMU Project, “Arm MPS2 and MPS3 boards,” QEMU System Emulation Documentation. Available: https://www.qemu.org/docs/master/system/arm/mps2.html. Accessed October 18, 2026.')
    if p._p.pPr is not None:p._p.remove(p._p.pPr)
    p._p.insert(0,deepcopy(ref_template))
    (root/"publication/array/postfreeze_edits.json").write_text(json.dumps(changes,indent=2,ensure_ascii=False)+"\n")
    assert len(abstract.text.split())<=250
    return dict(figures=11,tables=9,references=24,abstract_words=len(abstract.text.split()))
