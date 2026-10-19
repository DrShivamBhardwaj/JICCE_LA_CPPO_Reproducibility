"""Declare the authorized virtual benchmark scope without changing outcomes or gates."""
from pathlib import Path
from io import BytesIO
from lxml import etree
from docx import Document
import subprocess,json,hashlib,zipfile
R=Path(__file__).resolve().parents[1];P=R/'publication/array';B=R/'build/virtual_scope';B.mkdir(parents=True,exist_ok=True)
BASE='5d199b30209affa9d1fe67d54b070ad97f710b40'
raw=subprocess.check_output(['git','show',BASE+':publication/array/Array_Manuscript.docx'],cwd=R);old=Document(BytesIO(raw));d=Document(BytesIO(raw));changes=[]
for p in d.paragraphs:
 before=p.text
 if p.text.startswith('Forwarding silence can reflect'):
  p.text=p.text.replace('Physical RF, measured energy and production device protections remain unvalidated.','Physical RF, measured energy and production protections are outside this simulation/emulation evaluation and remain unvalidated.')
 elif p.text.startswith('The retained datasets therefore support virtual radio-state'):
  p.text='The evaluated scope ends at radio simulation and compiled-MCU emulation. Physical RF/current acquisition, protected device keys, authenticated physical provisioning, firmware authenticity and reset-safe MCU persistence are outside that scope and have not been validated. The benchmark records virtual endpoint behavior, including deadline failures; it does not certify deployment readiness. Native Cooja scheduling assumptions provide neither a serial CPU model nor physical calibration.'
 elif p.text.startswith('Exploratory reconstruction explains the censored'):
  p.text=p.text.replace('Physical RF, measured energy, protected device keys and reboot-persistent MCU state remain unvalidated.','Physical RF/current and production security lie outside this simulation/emulation benchmark; neither is validated by these results.')
 if before!=p.text:changes.append(dict(before=before,after=p.text))
assert len(changes)==3
math=lambda x:[etree.tostring(e) for e in x.element.xpath('.//m:oMath')]
assert math(old)==math(d) and len(math(d))==10
assert [t._tbl.xml for t in old.tables]==[t._tbl.xml for t in d.tables] and len(d.tables)==13
for p in old.paragraphs:
 if __import__('re').match(r'^\d{2}\.',p.text):
  q=next(q for q in d.paragraphs if q.text==p.text);assert etree.tostring(p._p)==etree.tostring(q._p)
assert len(d.inline_shapes)==16
(B/'baseline.docx').write_bytes(raw);d.save(P/'Array_Manuscript.docx')
(P/'Array_Manuscript.md').write_text('\n\n'.join(p.text for p in d.paragraphs)+'\n')
(P/'virtual_scope_edits.json').write_text(json.dumps(changes,indent=2)+'\n')
# Retain the complete historical-to-current audit mapping.
path=P/'followup_edits.json';mapping=json.loads(path.read_text());scope={c['before']:c['after'] for c in changes}
for item in mapping:item['after']=scope.get(item['after'],item['after'])
path.write_text(json.dumps(mapping,indent=2)+'\n')
(P/'virtual_scope_preservation_audit.json').write_text(json.dumps(dict(status='passed',source_commit=BASE,baseline_docx_sha256=hashlib.sha256(raw).hexdigest(),final_docx_sha256=hashlib.sha256((P/'Array_Manuscript.docx').read_bytes()).hexdigest(),changed_paragraphs=3,tables_preserved=13,figures_preserved=16,native_equations_preserved=10,deadline_acceptance_changed=False,physical_validation_claimed=False,production_security_claimed=False,user_authorized_scope='Option 1 simulation/emulation benchmark, 2026-10-19T07:24:15+05:30'),indent=2)+'\n')
print('Scope revised; every table, native equation and algorithm step preserved')
