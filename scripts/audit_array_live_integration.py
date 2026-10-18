"""Reconcile the new manuscript table and preserve the preceding evidence."""
from pathlib import Path
import argparse,json,hashlib,zipfile
from lxml import etree
from docx import Document
r=Path(__file__).resolve().parents[1];p=argparse.ArgumentParser();p.add_argument('--baseline',type=Path);a=p.parse_args()
file=r/'publication/array/Array_Manuscript.docx';d=Document(file);assert len(d.tables)==10 and len(d.inline_shapes)==11
runs=['outputs/cc2538_live/chain','outputs/cc2538_live/choice','outputs/cc2538_v2/chain_100','outputs/cc2538_v2/choice_100']
for i,folder in enumerate(runs):
 x=json.loads((r/folder/'verification.json').read_text());assert x['status']=='passed';delays=x['emulated_receipt_delays'];miss=sum(z['receipt_delay_ms']>[90,350,1500][z['traffic']] for z in delays)
 expected=[('Chain' if x['topology']=='chain' else 'Choice')+(' private' if x['expected_packets']==100 else ' public'),str(x['expected_packets']),f"{x['unique_sink_deliveries']} / {x['unique_source_verified_receipts']}",str(x['mcu_policy_decisions']),f"{min(z['receipt_delay_ms'] for z in delays)}–{max(z['receipt_delay_ms'] for z in delays)}",str(miss)]
 assert [c.text for c in d.tables[9].rows[i+1].cells]==expected
report=dict(status='passed',live_table_rows_reconciled=4,figures=11,tables=10,physical_validation=False,deadline_compliance=False)
if a.baseline:
 base=Document(a.baseline);assert len(base.tables)==9
 assert all(x._tbl.xml==y._tbl.xml for x,y in zip(base.tables,d.tables[:9]))
 def xmls(path):
  with zipfile.ZipFile(path) as z:
   tree=etree.fromstring(z.read('word/document.xml'))
   math=[etree.tostring(x,method='c14n') for x in tree.xpath('//*[local-name()="oMath"]')]
   media=sorted(hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/'))
  return math,media
 before,im0=xmls(a.baseline);after,im1=xmls(file);assert before==after and len(after)==10 and im0==im1
 report.update(preceding_nine_tables_unchanged=True,ten_native_equations_unchanged=True,preceding_eleven_images_unchanged=True,baseline_sha256=hashlib.sha256(a.baseline.read_bytes()).hexdigest())
(r/'publication/array/live_manuscript_audit.json').write_text(json.dumps(report,indent=2)+'\n');print(json.dumps(report,indent=2))
